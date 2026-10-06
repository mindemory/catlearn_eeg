"""Simple 2x2 design deck in the style of ~/Downloads/02_Mrugank.pptx: one idea per slide,
as little text as possible. Keeps that deck's master, layouts and embedded Nunito font,
drops its slides, and builds new ones from shapes and the real stimulus images.

Reads:
  model curves   ~/Documents/data/catlearn_eeg/kernel_model/2x2/matched_attention/predictions.json
                 (kernel_model/K02_matched_attention.py)
  stimuli        ~/Documents/data/catlearn_eeg/task_design/stimuli*/  (cropped copies are cached in
                 ~/Documents/data/catlearn_eeg/task_design/slide_assets/ on first run)

Needs python-pptx (pip install python-pptx); numpy and Pillow come with it or the env.

  python design_deck.py [out.pptx] [slide numbers to preview, e.g. 5,6]
"""

import json
import math
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

DATA = Path.home() / "Documents" / "data" / "catlearn_eeg"
TEMPLATE = Path.home() / "Downloads" / "02_Mrugank.pptx"
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path.home() / "Downloads" / "2x2 Design (simple).pptx"
ONLY = sys.argv[2] if len(sys.argv) > 2 else None          # e.g. "5" to preview one slide
STIM = DATA / "task_design"
ASSETS = STIM / "slide_assets"
PRED = json.loads((DATA / "kernel_model" / "2x2" / "matched_attention" / "predictions.json").read_text())

FONT = "Nunito SemiBold"
WHITE, GREY, DIM = "FFFFFF", "BFBFBF", "7F7F7F"
SCREEN = "8C8C8C"                     # mid-grey display, as in the task
A_COL, B_COL, AB_COL = "FFAB40", "A99BF0", "3CC9A0"
RULE_COLORS = {"A": A_COL, "B": B_COL, "AB": AB_COL}
F_COL, J_COL = "F2C14E", "4CC9F0"
GOOD, BAD = "3CC9A0", "FF6B6B"

FRACT = {"a1": "fractal34", "a2": "fractal58", "b1": "fractal17", "b2": "fractal66"}


def rgb(h):
    return RGBColor.from_string(h)


def make_assets():
    """Cropped, downscaled stimulus PNGs for the slides (transparent margins removed)."""
    import numpy as np
    from PIL import Image
    ASSETS.mkdir(parents=True, exist_ok=True)

    def crop_save(src, dst, size=360):
        if dst.exists():
            return
        im = Image.open(src).convert("RGBA")
        a = np.array(im)[..., 3]
        r, c = np.where(a.any(1))[0], np.where(a.any(0))[0]
        im = im.crop((c[0], r[0], c[-1] + 1, r[-1] + 1))
        side = max(im.size)
        canvas = Image.new("RGBA", (side, side), (0, 0, 0, 0))
        canvas.paste(im, ((side - im.width) // 2, (side - im.height) // 2))
        canvas.resize((size, size), Image.LANCZOS).save(dst)

    for k in [3, 9, 17, 24, 34, 39, 45, 51, 58, 62, 66, 71]:
        crop_save(STIM / "stimuli_fractals" / f"{k}.png", ASSETS / f"fractal{k}.png")
    for k in [1, 2, 6, 8, 12, 14]:
        crop_save(STIM / "stimuli_aliens" / f"alien{k:02d}.png", ASSETS / f"alien{k:02d}.png")
    for k in range(1, 7):
        crop_save(STIM / "stimuli" / "png" / f"noisepatch_{k:02d}.png", ASSETS / f"noise{k:02d}.png", size=256)


# ---------------------------------------------------------------- primitives

def text(slide, x, y, w, h, runs, size=18, color=WHITE, align="l", anchor="t", bold=False, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(0.02))
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "ctr": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(runs, str):
        runs = [(runs, {})]
    para = tf.paragraphs[0]
    for chunk, style in runs:
        for k, piece in enumerate(chunk.split("\n")):
            if k > 0:
                para = tf.add_paragraph()
            para.alignment = {"l": PP_ALIGN.LEFT, "ctr": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
            if not piece:
                continue
            r = para.add_run()
            r.text = piece
            f = r.font
            f.name = FONT
            f.size = Pt(style.get("size", size))
            f.bold = style.get("bold", bold)
            f.italic = style.get("italic", italic)
            f.color.rgb = rgb(style.get("color", color))
    return tb


def shape(slide, kind, x, y, w, h, fill=None, line=None, line_w=1.0, adj=None, dash=None):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if adj is not None:
        s.adjustments[0] = adj
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(line_w)
        if dash:
            s.line.dash_style = dash
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def label_in(s, txt, size=16, color=WHITE, bold=False):
    tf = s.text_frame
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(0))
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = txt
    r.font.name, r.font.size, r.font.bold = FONT, Pt(size), bold
    r.font.color.rgb = rgb(color)
    return s


def line(slide, x1, y1, x2, y2, color=WHITE, width=1.5, arrow=False, dash=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = dash
    if arrow:
        tail = etree.SubElement(c.line._get_or_add_ln(), qn("a:tailEnd"))
        tail.set("type", "triangle"); tail.set("w", "med"); tail.set("len", "med")
    return c


def polyline(slide, pts, color, width=2.0, dash=None):
    if max(y for _, y in pts) - min(y for _, y in pts) < 1e-4:
        return line(slide, pts[0][0], pts[0][1], pts[-1][0], pts[-1][1], color=color, width=width, dash=dash)
    fb = slide.shapes.build_freeform(Inches(pts[0][0]), Inches(pts[0][1]), scale=1.0)
    fb.add_line_segments([(Inches(x), Inches(y)) for x, y in pts[1:]], close=False)
    s = fb.convert_to_shape()
    s.fill.background()
    s.line.color.rgb = rgb(color)
    s.line.width = Pt(width)
    if dash:
        s.line.dash_style = dash
    s.shadow.inherit = False
    return s


def image(slide, name, cx, cy, size):
    return slide.shapes.add_picture(str(ASSETS / f"{name}.png"), Inches(cx - size / 2), Inches(cy - size / 2),
                                    Inches(size), Inches(size))


def fixation(slide, cx, cy, size=0.16, color=WHITE):
    line(slide, cx - size / 2, cy, cx + size / 2, cy, color=color, width=1.75)
    line(slide, cx, cy - size / 2, cx, cy + size / 2, color=color, width=1.75)


def screen(slide, x, y, w, h, left=None, right=None, item=0.42, extra=None):
    """Grey display with fixation and optional items left/right of it."""
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=SCREEN, adj=0.06)
    cx, cy = x + w / 2, y + h / 2
    fixation(slide, cx, cy, size=min(w, h) * 0.12)
    off = w * 0.3
    if left:
        image(slide, left, cx - off, cy, item)
    if right:
        image(slide, right, cx + off, cy, item)
    return cx, cy, off


def title(slide, txt):
    slide.shapes.title.text = txt
    for p in slide.shapes.title.text_frame.paragraphs:
        for r in p.runs:
            r.font.name = FONT
            r.font.size = Pt(28)
            r.font.color.rgb = rgb(WHITE)


def new_slide(prs, ttl=None):
    layout = prs.slide_layouts[4] if ttl else prs.slide_layouts[10]     # TITLE_ONLY / BLANK
    s = prs.slides.add_slide(layout)
    for ph in list(s.placeholders):
        if ph.placeholder_format.type != 1:          # keep only the title
            ph._element.getparent().remove(ph._element)
    if ttl:
        title(s, ttl)
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb("000000")
    return s


def pill(slide, x, y, w, h, rule, size=16):
    p = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=RULE_COLORS[rule], adj=0.5)
    label_in(p, rule, size=size, color="111111", bold=True)


def rule_grid(slide, x, y, cell, rule, keys=True):
    for la in (0, 1):
        for lb in (0, 1):
            cat = {"A": la, "B": lb, "AB": la ^ lb}[rule]
            c = shape(slide, MSO_SHAPE.RECTANGLE, x + lb * cell, y + la * cell, cell, cell,
                      fill=(F_COL, J_COL)[cat], line="000000", line_w=2)
            if keys:
                label_in(c, "FJ"[cat], size=int(cell * 30), color="111111", bold=True)


def plot(slide, x, y, w, h, series, ylim, zero=None, ylab=None):
    """series: [(values, color, dash)] drawn over x in [0, 1]."""
    y0, y1 = ylim
    line(slide, x, y + h, x + w, y + h, color=DIM, width=1)
    line(slide, x, y, x, y + h, color=DIM, width=1)
    if zero is not None:
        yz = y + h * (1 - (zero - y0) / (y1 - y0))
        line(slide, x, yz, x + w, yz, color="404040", width=1, dash=MSO_LINE_DASH_STYLE.DASH)
    for vals, color, dash in series:
        n = len(vals)
        pts = [(x + w * i / (n - 1), y + h * (1 - (min(max(v, y0), y1) - y0) / (y1 - y0))) for i, v in enumerate(vals)]
        polyline(slide, pts, color, 2.5 if not dash else 2.0, dash)
    if ylab:
        for v, lab, col in ylab:
            yv = y + h * (1 - (v - y0) / (y1 - y0))
            text(slide, x - 0.75, yv - 0.15, 0.7, 0.3, lab, size=12, color=col, align="r", anchor="ctr")


def smooth(vals, w=4):
    out = []
    for i in range(len(vals)):
        lo, hi = max(0, i - w // 2), min(len(vals), i + w - w // 2)
        out.append(sum(vals[lo:hi]) / (hi - lo))
    return out


def block(key, learner, b, what="acc"):
    learner = {"selective": "focused", "distributed": "spread"}.get(learner, learner)
    v = PRED["seqs"][key][learner]["acc" if what == "acc" else "attn_index"][40 * b:40 * (b + 1)]
    return smooth(v) if what == "acc" else v


# ---------------------------------------------------------------- slides

def s_title(prs):
    s = prs.slides.add_slide(prs.slide_layouts[0])
    for ph in list(s.placeholders):
        if ph.placeholder_format.type not in (3, 4):
            ph._element.getparent().remove(ph._element)
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb("000000")
    for ph in s.placeholders:
        ph.text = ("Does attention follow what we learn?" if ph.placeholder_format.type == 3
                   else "A 2×2 category-learning EEG study\nMrugank Dake · Oct 1, 2026")
        for p in ph.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.name = FONT
                r.font.size = Pt(34 if ph.placeholder_format.type == 3 else 18)
                r.font.color.rgb = rgb(WHITE)
    for i, k in enumerate(["fractal34", "fractal17", "fractal58", "fractal66"]):
        image(s, k, 3.35 + i * 1.1, 4.55, 0.75)


def s_question(prs):
    s = new_slide(prs)
    text(s, 0.8, 0.9, 8.4, 1.2, "When we learn which feature matters…", size=30, align="ctr")
    cx, cy, off = screen(s, 3.0, 2.0, 4.0, 2.25, left=FRACT["a1"], right=FRACT["b1"], item=0.8)
    text(s, cx - off - 0.6, 4.35, 1.2, 0.4, "?", size=28, color=A_COL, align="ctr")
    text(s, cx + off - 0.6, 4.35, 1.2, 0.4, "?", size=28, color=B_COL, align="ctr")
    text(s, 0.8, 4.75, 8.4, 0.6, "…does attention go there?", size=30, align="ctr")


def s_display(prs):
    s = new_slide(prs, "Two pictures, eyes on the cross")
    cx, cy, off = screen(s, 1.5, 1.35, 7.0, 3.7, left=FRACT["a1"], right=FRACT["b1"], item=1.2)
    text(s, cx - off - 0.6, 1.5, 1.2, 0.45, "A", size=24, color=A_COL, align="ctr", bold=True)
    text(s, cx + off - 0.6, 1.5, 1.2, 0.45, "B", size=24, color=B_COL, align="ctr", bold=True)
    line(s, cx, cy + 0.95, cx - off, cy + 0.95, color=WHITE, width=1.5, arrow=True)
    text(s, cx - off, cy + 1.0, off, 0.35, "6°", size=16, align="ctr")


def s_trial(prs):
    s = new_slide(prs, "One trial")
    w, h = 2.6, 1.46
    steps = [("Fixation", "1 s", None, None, None),
             ("Pictures flicker", "2 s", FRACT["a1"], FRACT["b1"], None),
             ("Press F or J", "≤ 2 s", None, None, None),
             ("Feedback", "0.5 s", None, None, "Correct")]
    for i, (name, dur, l, r, fb) in enumerate(steps):
        x, y = 0.7 + i * 1.45, 1.25 + i * 0.95
        cx, cy, off = screen(s, x, y, w, h, left=l, right=r, item=0.48)
        if fb:
            text(s, cx - 0.8, cy + 0.16, 1.6, 0.4, fb, size=17, color="2EE07A", align="ctr", bold=True)
        text(s, x + w + 0.1, y + 0.05, 2.6, 0.7, [(name + "\n", {"size": 18}), (dur, {"size": 15, "color": GREY})])
    line(s, 0.5, 2.95, 4.25, 5.42, color=WHITE, width=1.5, arrow=True)       # below-left of the cascade


def s_compounds(prs):
    s = new_slide(prs, "Four combinations")
    w, h = 2.6, 1.3
    for la in (0, 1):
        for lb in (0, 1):
            x, y = 2.6 + lb * 3.0, 1.45 + la * 1.75
            screen(s, x, y, w, h, left=FRACT[f"a{la + 1}"], right=FRACT[f"b{lb + 1}"], item=0.6)
    text(s, 1.4, 1.75, 1.0, 0.6, "a1", size=20, color=A_COL, align="r")
    text(s, 1.4, 3.5, 1.0, 0.6, "a2", size=20, color=A_COL, align="r")
    text(s, 2.6, 4.6, 2.6, 0.4, "b1", size=20, color=B_COL, align="ctr")
    text(s, 5.6, 4.6, 2.6, 0.4, "b2", size=20, color=B_COL, align="ctr")


def s_rules(prs):
    s = new_slide(prs, "Three rules")
    cell = 0.85
    labels = {"A": "Only A matters", "B": "Only B matters", "AB": "Both matter"}
    for i, rule in enumerate(("A", "B", "AB")):
        x = 1.0 + i * 3.0
        text(s, x - 0.4, 1.35, 2 * cell + 0.8, 0.45, rule, size=24, color=RULE_COLORS[rule], align="ctr", bold=True)
        rule_grid(s, x, 1.95, cell, rule)
        text(s, x - 0.5, 4.05, 2 * cell + 1.0, 0.4, labels[rule], size=16, align="ctr")
    text(s, 0.4, 2.1, 0.55, 0.5, "a1", size=14, color=A_COL, align="r")
    text(s, 0.4, 2.1 + cell, 0.55, 0.5, "a2", size=14, color=A_COL, align="r")
    text(s, 1.0, 1.95 + 2 * cell + 0.02, cell, 0.3, "b1", size=14, color=B_COL, align="ctr")
    text(s, 1.0 + cell, 1.95 + 2 * cell + 0.02, cell, 0.3, "b2", size=14, color=B_COL, align="ctr")
    text(s, 0.8, 4.75, 8.4, 0.5, "AB is XOR: neither picture alone tells you the answer", size=16, color=GREY, align="ctr")


def s_blocks(prs):
    s = new_slide(prs, "Three blocks, new pictures each time")
    seqs = [("A", "A", "AB"), ("A", "B", "AB"), ("AB", "AB", "A")]
    for b in range(3):
        text(s, 2.4 + b * 2.0, 1.35, 1.4, 0.35, f"block {b + 1}", size=14, color=GREY, align="ctr")
    for r, seq in enumerate(seqs):
        y = 1.85 + r * 0.85
        for b, rule in enumerate(seq):
            x = 2.4 + b * 2.0
            pill(s, x, y, 1.4, 0.55, rule, size=20)
            if b < 2:
                line(s, x + 1.48, y + 0.275, x + 1.92, y + 0.275, color=GREY, width=1.5, arrow=True)
    text(s, 0.8, 4.6, 8.4, 0.5, "40 trials per block · the rule changes without warning", size=16, color=GREY, align="ctr")


def spectrum(slide, x, y, w, h, big):
    """Cartoon EEG spectrum with peaks at 12 and 15 Hz; big = which one is taller."""
    f0, f1 = 8.0, 19.0
    pts = []
    for i in range(121):
        f = f0 + (f1 - f0) * i / 120
        v = 0.08 + 0.05 * math.exp(-(f - 10) ** 2 / 2)
        for fp, amp in ((12, 0.9 if big == 12 else 0.4), (15, 0.9 if big == 15 else 0.4)):
            v += amp * math.exp(-(f - fp) ** 2 / (2 * 0.12 ** 2))
        pts.append((x + w * i / 120, y + h * (1 - v)))
    line(slide, x, y + h, x + w, y + h, color=DIM, width=1)
    polyline(slide, pts, WHITE, 2.0)
    for fp, col in ((12, A_COL), (15, B_COL)):
        xf = x + w * (fp - f0) / (f1 - f0)
        text(slide, xf - 0.4, y + h + 0.05, 0.8, 0.3, f"{fp} Hz", size=12, color=col, align="ctr")


def s_tagging(prs):
    s = new_slide(prs, "Each picture flickers at its own rate")
    cx, cy, off = screen(s, 0.6, 1.5, 4.0, 2.2, left=FRACT["a1"], right=FRACT["b1"], item=0.8)
    text(s, cx - off - 0.6, 3.75, 1.2, 0.35, "12 Hz", size=16, color=A_COL, align="ctr")
    text(s, cx + off - 0.6, 3.75, 1.2, 0.35, "15 Hz", size=16, color=B_COL, align="ctr")
    text(s, 5.2, 1.3, 2.0, 0.35, "attend A", size=16, color=A_COL, align="ctr")
    spectrum(s, 5.2, 1.75, 2.0, 1.1, big=12)
    text(s, 7.5, 1.3, 2.0, 0.35, "attend B", size=16, color=B_COL, align="ctr")
    spectrum(s, 7.5, 1.75, 2.0, 1.1, big=15)
    text(s, 5.2, 3.45, 4.3, 0.4, "EEG power at each rate", size=14, color=GREY, align="ctr")
    text(s, 0.8, 4.55, 8.4, 0.5, "The bigger peak shows where attention is", size=18, align="ctr")


FOCUSED = (A_COL, None)
SPREAD = (WHITE, MSO_LINE_DASH_STYLE.DASH)


def legend(slide, x, y):
    polyline(slide, [(x, y + 0.17), (x + 0.45, y + 0.17)], A_COL, 2.5)
    text(slide, x + 0.55, y, 2.6, 0.34, "Focused on A", size=14, color=A_COL, anchor="ctr")
    polyline(slide, [(x + 2.4, y + 0.17), (x + 2.85, y + 0.17)], WHITE, 2.0, MSO_LINE_DASH_STYLE.DASH)
    text(slide, x + 2.95, y, 2.6, 0.34, "Spread over A and B", size=14, anchor="ctr")


def s_same_learning(prs):
    s = new_slide(prs, "Same learning, different attention")
    k = "A-A-AB"
    text(s, 1.2, 1.3, 3.4, 0.35, "Accuracy", size=18, align="ctr")
    plot(s, 1.2, 1.8, 3.4, 2.2, [(block(k, "selective", 0), *FOCUSED), (block(k, "distributed", 0), *SPREAD)],
         (0.45, 1.0), zero=0.5, ylab=[(1.0, "100%", GREY), (0.5, "50%", GREY)])
    text(s, 5.8, 1.3, 3.4, 0.35, "Attention", size=18, align="ctr")
    plot(s, 5.8, 1.8, 3.4, 2.2, [(block(k, "selective", 0, "att"), *FOCUSED), (block(k, "distributed", 0, "att"), *SPREAD)],
         (-1.0, 1.0), zero=0.0, ylab=[(1.0, "A", A_COL), (0.0, "even", GREY), (-1.0, "B", B_COL)])
    text(s, 1.2, 4.05, 3.4, 0.3, "trials in an A block", size=12, color=GREY, align="ctr")
    text(s, 5.8, 4.05, 3.4, 0.3, "trials in an A block", size=12, color=GREY, align="ctr")
    legend(s, 2.3, 4.6)


def s_after_switch(prs):
    s = new_slide(prs, "The difference shows when the rule changes")
    panels = [("Same rule again (A → A)", "A-A-AB"), ("Rule switches (A → B)", "A-B-AB")]
    for i, (lab, k) in enumerate(panels):
        x = 1.2 + i * 4.6
        text(s, x - 0.3, 1.3, 4.0, 0.35, lab, size=18, align="ctr")
        plot(s, x, 1.8, 3.4, 2.2, [(block(k, "selective", 1), *FOCUSED), (block(k, "distributed", 1), *SPREAD)],
             (0.45, 1.0), zero=0.5, ylab=[(1.0, "100%", GREY), (0.5, "50%", GREY)] if i == 0 else None)
        text(s, x, 4.05, 3.4, 0.3, "trials in block 2", size=12, color=GREY, align="ctr")
    legend(s, 2.3, 4.6)


def s_why_not_222(prs):
    s = new_slide(prs, "Why not 2×2×2 first?")
    # three items around fixation: left, right, above
    x, y, w, h = 0.6, 1.4, 3.6, 2.6
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=SCREEN, adj=0.06)
    cx, cy = x + w / 2, y + h / 2 + 0.15
    fixation(s, cx, cy, 0.25)
    for name, dx, dy, lab, col in (("fractal34", -1.15, 0, "A", A_COL), ("fractal17", 1.15, 0, "B", B_COL),
                                   ("fractal45", 0, -0.85, "C", AB_COL)):
        image(s, name, cx + dx, cy + dy, 0.6)
    text(s, x, y + h + 0.1, w, 0.35, "3 pictures · 8 combinations · 6 rule types", size=13, color=GREY, align="ctr")
    items = [
        ("Three flicker rates", "harmonics collide (7.5 × 2 = 15 Hz)"),
        ("No symmetric layout", "one picture sits above or below"),
        ("Attention split 3 ways", "weaker tagging signal"),
        ("8 combinations", "far more trials to learn and to average"),
        ("Hard rules often never learned", "nothing to compare attention against"),
        ("More pull to move the eyes", "eye artifacts in the EEG"),
    ]
    for i, (head, sub) in enumerate(items):
        yy = 1.3 + i * 0.6
        text(s, 4.7, yy, 5.0, 0.6, [(head + "\n", {"size": 16}), (sub, {"size": 12, "color": GREY})])


def s_stimuli(prs):
    s = new_slide(prs, "Which pictures?")
    sets = [("Noise patches", ["noise01", "noise02", "noise03", "noise04"],
             [("✓", "matched contrast"), ("✓", "hard to name"), ("✗", "hard to tell apart off-center")]),
            ("Aliens", ["alien01", "alien06", "alien08", "alien14"],
             [("✓", "easy to tell apart"), ("✗", "easy to name"), ("✗", "only 19 (12 per person)")]),
            ("Fractals", ["fractal34", "fractal58", "fractal17", "fractal66"],
             [("✓", "easy to tell apart"), ("✓", "72 to choose from"), ("✗", "brightness and color vary")])]
    for i, (name, imgs, notes) in enumerate(sets):
        x = 0.5 + i * 3.1
        text(s, x, 1.25, 2.9, 0.4, name, size=20, align="ctr")
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x + 0.25, 1.75, 2.4, 1.45, fill=SCREEN, adj=0.06)
        for j, im in enumerate(imgs):
            image(s, im, x + 0.85 + (j % 2) * 1.2, 2.1 + (j // 2) * 0.73, 0.62)
        for j, (mark, note) in enumerate(notes):
            yy = 3.35 + j * 0.38
            text(s, x + 0.25, yy, 0.3, 0.35, mark, size=15, color=GOOD if mark == "✓" else BAD, bold=True)
            text(s, x + 0.55, yy, 2.4, 0.35, note, size=13)
    text(s, 0.6, 4.6, 8.8, 0.5, "In our model, XOR was learned with aliens and fractals, not with noise patches",
         size=13, color=GREY, align="ctr")


def s_predictions(prs):
    s = new_slide(prs, "What we expect")
    items = [("1", "XOR takes longer than A or B"),
             ("2", "Attention moves to the picture that matters"),
             ("3", "Focused attention helps when the rule repeats,\nhurts when it changes"),
             ("4", "Attention before the response predicts accuracy")]
    for i, (n, t) in enumerate(items):
        y = 1.5 + i * 0.95
        c = shape(s, MSO_SHAPE.OVAL, 1.0, y, 0.5, 0.5, fill=A_COL)
        label_in(c, n, size=18, color="111111", bold=True)
        text(s, 1.75, y - 0.2, 7.5, 0.9, t, size=20, anchor="ctr")


SLIDES = [s_title, s_question, s_display, s_trial, s_compounds, s_rules, s_blocks, s_tagging,
          s_same_learning, s_after_switch, s_why_not_222, s_stimuli, s_predictions]


def main():
    make_assets()
    prs = Presentation(str(TEMPLATE))
    ids = prs.slides._sldIdLst
    for sld in list(ids):
        prs.part.drop_rel(sld.get(qn("r:id")))
        ids.remove(sld)
    for i, make in enumerate(SLIDES, 1):
        if ONLY is None or str(i) in ONLY.split(","):
            make(prs)
    prs.save(str(OUT))
    print(f"saved {OUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
