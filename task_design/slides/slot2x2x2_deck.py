"""2x2x2 slot-machine study deck (noise patches), in the style of the 4x4 deck.

Describes the online task in catlearn_2x2x2_prolific, the noise stimuli, their DreamSim /
DINO similarity analysis and the observer-model predictions for d' (kernel_model/
K06_noise_dprime.py). One idea per slide, little text. Timing, coins, criteria, rules and
conditions are read from the task's own src/config.js, so the slides match what
participants actually see.

Builds on the 4x4 deck (~/Downloads/4x4 Slot Machine Study.pptx: same master, layouts and
embedded font; its slides are dropped) and reuses its drawing helpers (slot4x4_deck.py).

  python slot2x2x2_deck.py [out.pptx] [slide numbers to preview, e.g. 2,3]
"""

import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

import design_deck as dd
import slot4x4_deck as s4
from design_deck import A_COL, B_COL, GREY, WHITE, image, label_in, line, shape, text
from slot4x4_deck import (COIN_COL, GOLD, KNOB, LOSE_COL, WIN_COL, arrow, bullseye, coin_bar, key_cap, new_slide,
                          oval, pill)

REPO = Path(__file__).resolve().parents[2]
TASK = REPO / "catlearn_2x2x2_prolific"
CONFIG = TASK / "src" / "config.js"
TEMPLATE = Path.home() / "Downloads" / "4x4 Slot Machine Study.pptx"
DATA = Path.home() / "Documents" / "data" / "catlearn_eeg"
SIMILARITY = DATA / "task_design" / "noise_similarity"
PREDICTIONS = DATA / "kernel_model" / "2x2x2" / "noise_dprime"
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path.home() / "Downloads" / "2x2x2 Slot Machine Study.pptx"
ONLY = sys.argv[2] if len(sys.argv) > 2 else None

SCREEN = "808080"                     # the task's background (LAYOUT.background); the patches' mean grey
C_COL = "3CC9A0"
POS_COL = {"A": A_COL, "B": B_COL, "C": C_COL}
STUDY_COLORS = ["4CC9F0", "2EC4B6", "FFA630", "C77DFF"]   # per study alpha, as in K06's figures
WIDE_ALPHAS = [0.5, 0.75, 1.0, 1.25, 1.5, 1.75, 2.0, 2.25, 2.5, 2.75, 3.0]


# ---------------------------------------------------------------- task settings from config.js

def read_config():
    js = CONFIG.read_text()

    def block(name):
        m = re.search(rf"export const {name} = (\{{.*?\n\}});", js, re.S)
        if not m:
            raise ValueError(f"{name} not found in {CONFIG}")
        return m.group(1)

    def num(src, key):
        m = re.search(rf"\b{key}:\s*(-?[\d.]+)", src)
        if not m:
            raise ValueError(f"{key} not found")
        return float(m.group(1))

    timing, reward = block("TIMING"), block("REWARD")
    practice, test = block("PRACTICE_CRITERION"), block("TEST_CRITERION")
    rules = {}
    for name, typ, levels, labels in re.findall(
            r"(\w+):\s*\{\s*type:\s*'([^']+)',\s*levels:\s*\[([^\]]+)\],\s*labels:\s*\[([^\]]+)\]", block("RULES")):
        rules[name] = {"type": typ, "levels": [int(v) for v in levels.split(",")],
                       "labels": [int(v) for v in labels.split(",")]}
    blocks = re.search(r"export const BLOCKS = \[(.*?)\n\];", js, re.S).group(1)
    reps = {ph: [int(r) for r in re.findall(rf"phase:\s*'{ph}'[^}}]*?reps:\s*(\d+)", blocks)]
            for ph in ("practice", "test")}
    test_rule = re.search(r"phase:\s*'test',\s*rule:\s*'(\w+)'", blocks).group(1)
    positions = {p: (float(x), float(y)) for p, x, y in
                 re.findall(r"(\w):\s*\[(-?[\d.]+),\s*(-?[\d.]+)\]", re.search(r"positions:\s*\{([^}]*)\}", js).group(1))}
    m = re.search(r"machine:\s*\{\s*center:\s*\[(-?[\d.]+),\s*(-?[\d.]+)\],\s*size:\s*\[([\d.]+),\s*([\d.]+)\]", js)
    cfg = {
        "wait": num(timing, "wait") / 1000, "pull": num(timing, "pull") / 1000,
        "response": num(timing, "response") / 1000, "feedback": num(timing, "feedback") / 1000,
        "correct": int(num(reward, "correct")), "wrong": int(num(reward, "wrong")), "late": int(num(reward, "late")),
        "coins_per_dollar": int(num(reward, "coinsPerDollar")),
        "window": int(num(practice, "window")), "min_correct": int(num(practice, "minCorrect")),
        "test_window": int(num(test, "window")), "test_min_correct": int(num(test, "minCorrect")),
        "credit": "creditRemaining: true" in test,
        "rules": rules, "test_rule": test_rule,
        "practice_max": [r * 4 for r in reps["practice"]], "test_trials": [r * 8 for r in reps["test"]],
        "alphas": re.findall(r"'([\d.]+)'", re.search(r"alphas:\s*\[([^\]]*)\]", block("STIMULI")).group(1)),
        "positions": positions, "symbol": num(block("LAYOUT"), "fractalSize"),
        "machine_center": (float(m.group(1)), float(m.group(2))), "machine_size": (float(m.group(3)), float(m.group(4))),
        "hint_y": num(block("LAYOUT"), "hintY"),
        "assumed_cm": num(block("CALIBRATION"), "assumedDistanceMm") / 10,
    }
    cfg["dist_range_cm"] = [int(v) / 10 for v in re.search(r"plausibleDistanceMm:\s*\[(\d+),\s*(\d+)\]", js).groups()]
    return cfg


CFG = read_config()


# ---------------------------------------------------------------- slide assets

def noise_target(alpha):
    import json
    return json.loads((TASK / "stimuli" / "noise" / f"a{alpha}" / "groups.json").read_text())["target_distance"]


def make_assets():
    """Noise patches (the task's own, and newly generated for the wide alpha range), the
    making-of steps, practice fractals and crops of the similarity figures."""
    import numpy as np
    from PIL import Image
    sys.path.insert(0, str(REPO / "task_design"))
    from noise_patches import circular_aperture, make_noise_patch
    A = dd.ASSETS
    A.mkdir(parents=True, exist_ok=True)

    def rgba(field, mask):
        grey = np.round(255 * (0.5 + 0.5 * field)).astype(np.uint8)
        return Image.fromarray(np.dstack([grey, grey, grey, np.round(255 * mask).astype(np.uint8)]))

    for a in CFG["alphas"]:                                  # the task's patches: files 1-8 = pairs 1-4
        for k in range(1, 9):
            dst = A / f"n3_a{a}_{k}.png"
            if not dst.exists():
                Image.open(TASK / "stimuli" / "noise" / f"a{a}" / f"{k}.png").convert("RGBA").resize(
                    (200, 200), Image.LANCZOS).save(dst)
    mask = circular_aperture(256)
    for a in WIDE_ALPHAS:                                    # wide range, generated here
        rng = np.random.default_rng([7, int(round(a * 100))])
        for k in range(1, 5):
            dst = A / f"n3_w{a:.2f}_{k}.png"
            field = make_noise_patch(256, rng, a)
            if not dst.exists():
                rgba(field, mask).resize((200, 200), Image.LANCZOS).save(dst)
    steps = {"n3_step_white": (0.0, np.ones_like(mask)), "n3_step_filtered": (2.0, np.ones_like(mask)),
             "n3_step_aperture": (2.0, mask)}
    for name, (a, m) in steps.items():                       # same white noise underneath every step
        dst = A / f"{name}.png"
        if not dst.exists():
            rgba(make_noise_patch(256, np.random.default_rng(11), a), m).save(dst)
    for k in range(1, 5):                                    # practice fractals: 2x2x2 set, pairs 1 and 2
        dst = A / f"n3_f{k}.png"
        if dst.exists():
            continue
        im = Image.open(TASK / "stimuli" / "fractal_groups" / f"{k}.png").convert("RGBA")
        al = np.array(im)[..., 3]
        r, c = np.where(al.any(1))[0], np.where(al.any(0))[0]
        im = im.crop((c[0], r[0], c[-1] + 1, r[-1] + 1))
        side = max(im.size)
        canvas = Image.new("RGBA", (side, side), (0, 0, 0, 0))
        canvas.paste(im, ((side - im.width) // 2, (side - im.height) // 2))
        canvas.resize((300, 300), Image.LANCZOS).save(dst)
    crops = {"n3_sim_dreamsim": ("within_alpha", (0, 400, 462, 1245)),
             "n3_sim_dino": ("within_alpha", (468, 400, 922, 1245)),
             "n3_sim_local": ("within_alpha", (924, 400, 1385, 1245)),
             "n3_sim_pixel": ("within_alpha", (1386, 400, 1846, 1245)),
             "n3_between": ("between_alpha", (0, 50, 618, 616))}
    for name, (src, box) in crops.items():
        Image.open(SIMILARITY / f"{src}.png").crop(box).save(A / f"{name}.png")
    for name in ("baseline", "learning", "gain"):
        Image.open(PREDICTIONS / f"{name}.png").save(A / f"n3_pred_{name}.png")


def picture(slide, name, x, y, w=None, h=None):
    """An asset at (x, y) inches, sized by width or height; returns the picture."""
    return slide.shapes.add_picture(str(dd.ASSETS / f"{name}.png"), dd.Inches(x), dd.Inches(y),
                                    width=dd.Inches(w) if w else None, height=dd.Inches(h) if h else None)


def screen_row(slide, names, x, y, size, gap, fill=SCREEN):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x - 0.08, y - 0.08, len(names) * (size + gap) - gap + 0.16,
          size + 0.16, fill=fill, adj=0.12)
    for i, f in enumerate(names):
        image(slide, f, x + i * (size + gap) + size / 2, y + size / 2, size)


# ---------------------------------------------------------------- the machine

def example(kind="noise", alpha="2.00", pair_set=0):
    """Asset names for A, B, C: noise patches of one alpha (one patch of each of 3 pairs), or
    practice fractals for A and B."""
    if kind == "fractals":
        return {"A": "n3_f1", "B": "n3_f3"}
    k = [1, 3, 5] if pair_set == 0 else [2, 3, 6]
    return {p: f"n3_a{alpha}_{i}" for p, i in zip("ABC", k)}


def slot3(slide, cx, cy, u, files=None, outcome=None, pulled=False, bar=None, hint=None, screen=True, labels=False):
    """The task screen, u inches per degree, fixation at (cx, cy): grey display, gold machine
    (LAYOUT.machine) with plate and lever, symbols at LAYOUT.positions, bull's eye, optional
    feedback circles, coin bar (frac, label) and hint text."""
    at = lambda x, y: (cx + x * u, cy - y * u)
    if screen:
        shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, cx - 13 * u, cy - 12.6 * u, 26 * u, 21.6 * u, fill=SCREEN, adj=0.04)
    (mx, my), (w, h) = CFG["machine_center"], CFG["machine_size"]
    fx, fy = at(mx, my)
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, fx - w * u / 2, fy - h * u / 2, w * u, h * u, line=GOLD,
          line_w=max(0.75, 0.5 * u * 72), adj=0.06)
    plate = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, fx - 1.6 * u, fy - h * u / 2 - 0.45 * u, 3.2 * u, 0.9 * u,
                  fill=GOLD, adj=0.3)
    if u > 0.12:
        label_in(plate, "SLOTS", size=max(6, int(u * 18)), color="3A2A00", bold=True)
    px = fx + w * u / 2 + 0.9 * u
    arm_top = fy if pulled else fy - 3.5 * u
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, px - 0.15 * u, arm_top, 0.3 * u, 3.5 * u, fill=GOLD, adj=0.5)
    oval(slide, px, fy + 3.5 * u if pulled else fy - 3.5 * u, 1.0 * u, fill=KNOB)
    files = files or {}
    for p, f in files.items():
        image(slide, f, *at(*CFG["positions"][p]), CFG["symbol"] * u)
    if outcome:
        for p in files:
            oval(slide, *at(*CFG["positions"][p]), 5.4 * u, line_col=outcome, line_w=max(1.0, 0.3 * u * 72))
    bullseye(slide, cx, cy, u)
    if bar is not None:
        frac, lab = bar
        coin_bar(slide, cx - 1.0 * u, fy - h * u / 2 - 1.9 * u, 9 * u, frac, lab, size=max(7, int(u * 30)))
    if hint:
        hx, hy = at(0, CFG["hint_y"])
        text(slide, hx - 8 * u, hy - 0.4 * u, 16 * u, 0.8 * u, hint, size=max(7, int(u * 26)), color="E0E0E0",
             align="ctr", anchor="ctr")
    if labels:
        for p in labels if isinstance(labels, str) else "ABC":
            x, y = at(*CFG["positions"][p])
            dx = -3.4 * u if p in "AC" else 3.4 * u
            text(slide, x + dx - 0.3, y - 0.2, 0.6, 0.4, p, size=16, color=POS_COL[p], bold=True, align="ctr",
                 anchor="ctr")


# ---------------------------------------------------------------- slides: the task

def s_title(prs):
    s = prs.slides.add_slide(prs.slide_layouts[0])
    for ph in list(s.placeholders):
        if ph.placeholder_format.type not in (3, 4):
            ph._element.getparent().remove(ph._element)
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = dd.rgb("000000")
    for ph in s.placeholders:
        title = ph.placeholder_format.type == 3
        ph.text = "Slot machines with three symbols" if title else "A 2×2×2 online category-learning study with noise patches"
        for p in ph.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.name = dd.FONT
                r.font.size = Pt(34 if title else 18)
                r.font.color.rgb = dd.rgb(WHITE)
    names = [f"n3_a{a}_{k}" for a in CFG["alphas"] for k in (1, 3)]
    screen_row(s, names, 1.35, 4.25, 0.72, 0.18)


def s_machine(prs):
    s = new_slide(prs, "The slot machine: three symbols")
    u, cx, cy = 0.158, 3.6, 3.38
    slot3(s, cx, cy, u, files=example(), bar=(0.4, "1,240 coins"), labels=True)
    ax, ay = CFG["positions"]["C"]
    line(s, cx, cy, cx + ax * u, cy - (ay - 2.1) * u, color=WHITE, width=1.25, arrow=True)
    text(s, cx + 0.05, cy - 3.3 * u - 0.15, 0.5, 0.3, "6°", size=13)
    bx, by = CFG["positions"]["B"]
    text(s, cx + bx * u - 0.5, cy - (by - 2.1) * u, 1.0, 0.28, "4°", size=13, align="ctr")
    text(s, 6.25, 1.75, 3.4, 2.6, [("A, B, C\n", {"size": 18}),
                                   ("corners of an equilateral triangle, each 6° from the bull's eye\n\n",
                                    {"size": 13, "color": GREY}),
                                   ("4° symbols\n", {"size": 18}),
                                   ("noise patches in the test, fractals in practice", {"size": 13, "color": GREY})])


def s_rule(prs):
    s = new_slide(prs, "8 combinations, half of them win")
    labels = CFG["rules"][CFG["test_rule"]]["labels"]
    a = "2.00"
    heads = {"A": [f"n3_a{a}_1", f"n3_a{a}_2"], "B": [f"n3_a{a}_3", f"n3_a{a}_4"], "C": [f"n3_a{a}_5", f"n3_a{a}_6"]}
    cell, size = 0.6, 0.45
    for c in range(2):
        px = 1.0 + c * 3.35
        gx, gy = px + 1.05, 2.55
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, px, 1.35, 2.95, 2.55, fill="1A1A1A", adj=0.06)
        text(s, px + 0.75, 1.42, 0.6, 0.4, "C =", size=15, color=C_COL, bold=True, align="r", anchor="ctr")
        image(s, heads["C"][c], px + 1.65, 1.62, size)
        text(s, px + 0.2, gy - 0.6, 0.4, 0.4, "B", size=14, color=B_COL, bold=True, align="ctr", anchor="ctr")
        text(s, px + 0.05, gy + cell - 0.2, 0.3, 0.4, "A", size=14, color=A_COL, bold=True, anchor="ctr")
        for i in range(2):
            image(s, heads["A"][i], gx - 0.4, gy + (i + 0.5) * cell, size)
            image(s, heads["B"][i], gx + (i + 0.5) * cell, gy - 0.35, size)
            for j in range(2):
                win = labels[4 * i + 2 * j + c] == 1
                shape(s, MSO_SHAPE.RECTANGLE, gx + j * cell, gy + i * cell, cell - 0.05, cell - 0.05,
                      fill=WIN_COL if win else LOSE_COL)
    lx = 7.95
    shape(s, MSO_SHAPE.RECTANGLE, lx, 2.0, 0.32, 0.32, fill=WIN_COL)
    text(s, lx + 0.42, 1.97, 1.6, 0.4, "wins", size=15, anchor="ctr")
    shape(s, MSO_SHAPE.RECTANGLE, lx, 2.5, 0.32, 0.32, fill=LOSE_COL)
    text(s, lx + 0.42, 2.47, 1.6, 0.4, "loses", size=15, anchor="ctr")
    text(s, 0.6, 4.05, 8.8, 0.9,
         [(f"Type {CFG['rules'][CFG['test_rule']]['type']} · A + B + AC + BC · the same rule for everyone\n",
           {"size": 16}),
          ("any one or two positions alone get 75% right: all three are needed", {"size": 13, "color": GREY})],
         align="ctr")


def s_round(prs):
    s = new_slide(prs, "One round")
    u = 0.072
    steps = [
        ("Wait", f"SPACE, or {CFG['wait']:g} s", dict(hint="press SPACE to play")),
        ("Lever pull", f"{CFG['pull']:g} s", dict(pulled=True)),
        ("Will these win?", f"F / J, ≤ {CFG['response']:g} s", dict(files=example())),
        ("Feedback", f"{CFG['feedback']:g} s", dict(files=example(), outcome="2ECC40")),
    ]
    for i, (name, dur, kw) in enumerate(steps):
        cx, cy = 1.55 + i * 1.6, 1.95 + i * 0.85
        slot3(s, cx, cy, u, bar=(0.4, ""), **kw)
        text(s, cx + 1.0, cy - 0.95, 2.6, 0.7, [(name + "\n", {"size": 15}), (dur, {"size": 13, "color": GREY})])
    text(s, 6.4, 4.95, 3.4, 0.35, "then the next round", size=13, color=GREY, align="r")


def s_keys(prs):
    s = new_slide(prs, "Will these win? YES or NO")
    for i, yes in enumerate("FJ"):
        x = 1.4 + i * 4.2
        text(s, x - 0.2, 1.45, 3.2, 0.35, "Half the participants", size=14, color=GREY, align="ctr")
        for k, letter in enumerate("FJ"):
            key_cap(s, x + 0.55 + 1.1 * k, 2.0, letter, size=0.8, label="YES" if letter == yes else "NO",
                    label_col="2ECC40" if letter == yes else "FF6B6B")
    text(s, 0.6, 3.9, 8.8, 0.9, [("Left vs right hand is not tied to YES / NO\n", {"size": 16}),
                                 ("assigned at random per participant · upper or lower case both count",
                                  {"size": 13, "color": GREY})], align="ctr")


def s_feedback(prs):
    s = new_slide(prs, "Feedback: green or red circles")
    u = 0.112
    for i, (col, lab, sub, frac, bar_lab) in enumerate((
            ("2ECC40", f"+{CFG['correct']} coins", "correct", 0.45, "1,250 coins"),
            ("FF4136", f"−{abs(CFG['wrong'])} coins", "wrong or too slow", 0.42, "1,235 coins"))):
        cx = 2.55 + i * 4.9
        slot3(s, cx, 2.85, u, files=example(pair_set=1), outcome=col, bar=(frac, bar_lab))
        text(s, cx - 1.8, 4.05, 3.6, 0.6, [(lab + "\n", {"size": 18, "color": col}), (sub, {"size": 13, "color": GREY})],
             align="ctr")
    text(s, 0.6, 4.95, 8.8, 0.35, "a circle around every symbol · no text on the screen · the coin bar moves",
         size=13, color=GREY, align="ctr")


def s_practice(prs):
    s = new_slide(prs, "Practice: two 2×2 machines with fractals")
    r = CFG["rules"]
    items = [("Type I", "A or B, at random", r["x2_I_A"]["labels"]), ("XOR", "both matter", r["x2_XOR"]["labels"])]
    for i, (name, sub, labels) in enumerate(items):
        x = 4.6 + i * 2.7
        pill(s, x, 1.45, 1.9, 0.48, name, A_COL if i == 0 else "3CC9A0", size=16)
        s4.rule_grid(s, x + 0.4, 2.15, 0.55, labels, 2, gap=0.04)
        text(s, x - 0.3, 3.35, 2.5, 0.35, sub, size=13, color=GREY, align="ctr")
        if i == 0:
            arrow(s, x + 2.0, 1.69, x + 2.6)
    slot3(s, 2.15, 3.0, 0.1, files=example("fractals"), labels="AB")
    text(s, 0.6, 4.2, 8.8, 0.9,
         [(f"Each ends once {CFG['min_correct']} of the last {CFG['window']} answers are right\n", {"size": 17}),
          (f"at most {CFG['practice_max'][0]} rounds · A and B only · practice coins are not paid",
           {"size": 13, "color": GREY})], align="ctr")


def s_test(prs):
    s = new_slide(prs, "Test: two machines, same rule")
    for i, (name, pairs) in enumerate((("machine 1", 0), ("machine 2", 1))):
        cx = 2.6 + i * 4.8
        slot3(s, cx, 2.75, 0.085, files=example(pair_set=pairs) if i == 0 else
              {p: f"n3_a2.00_{k}" for p, k in zip("ABC", (7, 8, 2))})
        text(s, cx - 1.5, 3.85, 3.0, 0.35, name, size=16, align="ctr")
        if i == 0:
            arrow(s, cx + 1.4, 2.55, cx + 3.4)
            text(s, cx + 1.3, 2.05, 2.2, 0.4, "new patches", size=12, color=GREY, align="ctr")
    n, wnd, need = CFG["test_trials"][0], CFG["test_window"], CFG["test_min_correct"]
    text(s, 0.6, 4.3, 8.8, 0.9,
         [(f"Each: up to {n} rounds, ends once {need} of the last {wnd} are right\n", {"size": 17}),
          ("rounds left are paid as correct, so learning fast never costs bonus" if CFG["credit"] else "",
           {"size": 13, "color": GREY})], align="ctr")


def s_criterion(prs):
    s = new_slide(prs, f"Why {CFG['test_min_correct']} of {CFG['test_window']}?")
    text(s, 0.6, 1.3, 8.8, 0.4, "simulated, 300 rounds: share of participants who reach the criterion", size=14,
         color=GREY, align="ctr")
    rows = [("12 of 16", "91%", "100%", "100%"), ("18 of 24", "45%", "100%", "100%"),
            ("21 of 24", "1%", "30%", "100%"), ("22 of 24", "0%", "2%", "100%")]
    cols = [("criterion", 1.8), ("guessing", 1.7), ("knows only A", 1.9), ("knows the rule", 1.9)]
    x0, y0 = 1.35, 1.85
    x = x0
    for head, w in cols:
        text(s, x, y0, w, 0.35, head, size=13, color=GREY, align="ctr")
        x += w
    for r, vals in enumerate(rows):
        y = y0 + 0.42 + r * 0.48
        pick = vals[0] == f"{CFG['test_min_correct']} of {CFG['test_window']}"
        if pick:
            shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x0 - 0.1, y - 0.03, sum(w for _, w in cols) + 0.2, 0.46,
                  fill="2A2A2A", adj=0.3)
        x = x0
        for (head, w), v in zip(cols, vals):
            text(s, x, y, w, 0.4, v, size=17, color=WHITE if pick else GREY, bold=pick, align="ctr", anchor="ctr")
            x += w
    text(s, 0.6, 4.4, 8.8, 0.7, [("any one or two positions alone already give 75%: a 75% criterion "
                                  "can't tell them apart\n", {"size": 13}),
                                 ("“knows the rule” = 85–95% correct", {"size": 12, "color": GREY})], align="ctr")


def s_bonus(prs):
    s = new_slide(prs, "Coins become a bonus")
    n = sum(CFG["test_trials"])
    coins = lambda p: n * (p * CFG["correct"] + (1 - p) * CFG["wrong"])
    coin_bar(s, 4.6, 1.75, 5.0, 0.55, "1,240 coins", size=14)
    rows = [("guessing", f"{coins(0.5):,.0f}", 0.5), ("72% correct", f"~{round(coins(0.72), -1):,.0f}", 0.72),
            ("criterion reached", f"up to {coins(1.0):,.0f}", 1.0)]
    for i, (name, c, p) in enumerate(rows):
        y = 2.4 + i * 0.6
        text(s, 1.6, y, 2.8, 0.5, name, size=17, anchor="ctr")
        text(s, 4.2, y, 2.6, 0.5, f"{c} coins", size=17, color=COIN_COL, anchor="ctr", align="r")
        v = coins(p) / CFG["coins_per_dollar"]
        text(s, 7.0, y, 1.2, 0.5, f"${v:.2f}", size=17, color=GREY, anchor="ctr", align="r")
    text(s, 0.6, 4.4, 8.8, 0.8,
         [(f"Test machines only · {CFG['coins_per_dollar']:,} coins = $1 · never below $0\n", {"size": 14}),
          (f"over both machines ({n} rounds at most) · paid as a Prolific bonus on top of the base pay",
           {"size": 12, "color": GREY})], align="ctr")


def s_dva(prs):
    s = new_slide(prs, "Sizes in degrees of visual angle")
    text(s, 0.6, 1.3, 4.2, 0.35, "1 · match a real card", size=16, align="ctr")
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 1.55, 1.85, 2.2, 1.39, fill="DDDDDD", adj=0.06)
    shape(s, MSO_SHAPE.RECTANGLE, 3.5, 3.0, 0.25, 0.24, line="FF0000", line_w=2.5)
    text(s, 0.6, 3.35, 4.2, 0.35, "→ pixels per mm", size=14, color=GREY, align="ctr")
    text(s, 5.2, 1.3, 4.2, 0.35, "2 · find the blind spot", size=16, align="ctr")
    shape(s, MSO_SHAPE.RECTANGLE, 8.35, 2.35, 0.35, 0.35, fill="000000", line=WHITE, line_w=2)
    oval(s, 6.6, 2.52, 0.33, fill="FF2222", line_col=WHITE, line_w=1.5)
    line(s, 6.35, 2.95, 5.75, 2.95, color=GREY, width=1.5, arrow=True)
    text(s, 5.2, 3.05, 4.2, 0.6, [("cover the right eye, look AT the square,\n", {"size": 12}),
                                  ("SPACE when the ball vanishes → viewing distance", {"size": 12})],
         color=GREY, align="ctr")
    text(s, 0.6, 3.95, 8.8, 0.45, "pixels per degree = 2 · distance · tan(0.5°) · pixels per mm", size=16,
         align="ctr")
    lo, hi = CFG["dist_range_cm"]
    text(s, 0.6, 4.45, 8.8, 0.7,
         [(f"Outside {lo:g}–{hi:g} cm: measure again once, then assume {CFG['assumed_cm']:g} cm (flagged)\n", {}),
          ("too big for the window: everything shrinks by one factor (layout_scale, saved)", {})],
         size=12, color=GREY, align="ctr")


def s_online(prs):
    s = new_slide(prs, "Built for online data")
    items = [
        ("Desktop only", "browser check, window ≥ 1000 × 650 px"),
        ("Full screen", "exits are counted on every trial"),
        ("Tab switches", "counted on every trial"),
        ("Quiz", "instructions repeat until right (3 tries)"),
        ("Every trial saved", "symbols, key, RT, coins, px / degree, α"),
        ("Reproducible", "seed rebuilds the design · Prolific IDs on every row"),
    ]
    for i, (head, sub) in enumerate(items):
        col, row = i % 2, i // 2
        x, y = 0.9 + col * 4.4, 1.5 + row * 1.1
        shape(s, MSO_SHAPE.OVAL, x, y + 0.08, 0.22, 0.22, fill=A_COL)
        text(s, x + 0.4, y - 0.05, 3.9, 0.9, [(head + "\n", {"size": 16}), (sub, {"size": 12, "color": GREY})])


def session_minutes():
    """At most: ~0.8 s wait + pull + ~1.2 s answer + feedback per round, every test round,
    ~40 rounds per practice block, ~8 min for consent, screen setup and instructions."""
    per_round = 0.8 + CFG["pull"] + 1.2 + CFG["feedback"]
    return round((sum(CFG["test_trials"]) * per_round + 2 * 40 * per_round) / 60 + 8)


def s_session(prs):
    s = new_slide(prs, "The session")
    steps = [("Consent", None), ("Screen setup", "~2 min"), ("Instructions + quiz", None),
             ("Practice", "two 2×2 machines"), ("Test", f"2 × ≤ {CFG['test_trials'][0]}"), ("Bonus → Prolific", None)]
    x = 0.45
    widths = [1.15, 1.45, 1.75, 1.15, 1.15, 1.6]
    for (name, sub), w in zip(steps, widths):
        pill(s, x, 2.15, w, 0.6, name, "3A3A3A", size=12, color=WHITE)
        if sub:
            text(s, x, 2.85, w, 0.3, sub, size=11, color=GREY, align="ctr")
        x += w + 0.13
    text(s, 0.6, 3.9, 8.8, 0.5, f"at most about {session_minutes()} minutes", size=20, align="ctr")
    text(s, 0.6, 4.45, 8.8, 0.35, "shorter for anyone who reaches the criterion · rest between blocks", size=13,
         color=GREY, align="ctr")


# ---------------------------------------------------------------- slides: the noise patches

def s_alphas(prs):
    s = new_slide(prs, "Noise patches: one α per participant")
    size, gap = 0.62, 0.14
    for r, (a, col) in enumerate(zip(CFG["alphas"], STUDY_COLORS)):
        y = 1.32 + r * 0.86
        text(s, 1.0, y + 0.1, 1.1, 0.45, f"α {a}", size=17, color=col, bold=True, align="r", anchor="ctr")
        screen_row(s, [f"n3_a{a}_{k}" for k in (1, 3, 5, 7)], 2.35, y, size, gap)
        text(s, 5.6, y + 0.1, 3.6, 0.45, f"pairs {noise_target(a):.3f} apart", size=14, color=GREY, anchor="ctr")
    text(s, 0.6, 4.85, 8.8, 0.4, "larger α = smoother blobs = easier to tell apart · distance = DreamSim; "
         "fractal pairs: 0.256", size=12, color=GREY, align="ctr")


def s_alpha_range(prs):
    s = new_slide(prs, "α from 0.5 to 3.0")
    size, gap, rows = 0.44, 0.07, 6
    study = {float(a): c for a, c in zip(CFG["alphas"], STUDY_COLORS)}
    for i, a in enumerate(WIDE_ALPHAS):
        col, r = divmod(i, rows)
        x, y = 1.35 + col * 4.45, 1.3 + r * 0.6
        hit = study.get(a)
        text(s, x - 1.05, y + 0.02, 0.9, 0.4, f"α {a:g}", size=14, color=hit or GREY, bold=bool(hit), align="r",
             anchor="ctr")
        screen_row(s, [f"n3_w{a:.2f}_{k}" for k in range(1, 5)], x, y, size, gap)
        if hit:
            text(s, x + 4 * (size + gap) + 0.05, y + 0.02, 0.8, 0.4, "study", size=11, color=hit, anchor="ctr")
    text(s, 0.6, 4.95, 8.8, 0.35, "0 = white noise · ~1 = natural images · 3 = a few soft blobs", size=12,
         color=GREY, align="ctr")


def s_noise_how(prs):
    s = new_slide(prs, "How each patch is made")
    steps = [("n3_step_white", "white noise"), ("n3_step_filtered", "amplitude ∝ 1 / f^α"),
             ("n3_step_aperture", "soft circular aperture")]
    for i, (name, lab) in enumerate(steps):
        cx = 1.85 + i * 3.15
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, cx - 1.05, 1.5, 2.1, 2.1, fill=SCREEN, adj=0.06)
        image(s, name, cx, 2.55, 1.8)
        text(s, cx - 1.4, 3.7, 2.8, 0.35, lab, size=15, align="ctr")
        if i < 2:
            arrow(s, cx + 1.15, 2.55, cx + 2.0)
    text(s, 0.6, 4.35, 8.8, 0.8, [("same RMS contrast (0.3) at every α · mean = the background grey\n", {"size": 14}),
                                  ("example: α 2.0 · 4° across in the task", {"size": 12, "color": GREY})],
         align="ctr")


def s_noise_pairs(prs):
    s = new_slide(prs, "Equally distinct pairs within each α")
    ds = [noise_target(a) for a in CFG["alphas"]]
    top = 0.256
    x0, y_base, h = 1.4, 4.1, 2.5
    for i, (a, d, col) in enumerate(zip(CFG["alphas"], ds, STUDY_COLORS)):
        x = x0 + i * 1.75
        bh = h * d / top
        shape(s, MSO_SHAPE.RECTANGLE, x, y_base - bh, 0.9, bh, fill=col)
        text(s, x - 0.3, y_base - bh - 0.32, 1.5, 0.3, f"{d:.3f}", size=13, color=col, align="ctr")
        text(s, x - 0.3, y_base + 0.05, 1.5, 0.3, f"α {a}", size=13, color=col, bold=True, align="ctr")
        screen_row(s, [f"n3_a{a}_1", f"n3_a{a}_2"], x - 0.02, y_base - bh - 0.95, 0.42, 0.06)
    line(s, x0 - 0.3, y_base - h, x0 + 7.0, y_base - h, color="FFA630", width=1.25, dash=MSO_LINE_DASH_STYLE.DASH)
    text(s, x0 + 5.5, y_base - h - 0.33, 1.8, 0.3, "fractal pairs 0.256", size=12, color="FFA630")
    text(s, 0.6, 4.6, 8.8, 0.7, [("30 pairs per α, each within ±0.0013 of that α's median distance (DreamSim)\n",
                                  {"size": 13}),
                                 ("a participant sees 6 pairs: every position and machine is equally hard",
                                  {"size": 12, "color": GREY})], align="ctr")


def s_sim_within(prs):
    s = new_slide(prs, "DreamSim and DINO: noise pairs stay below fractals")
    for i, (name, lab) in enumerate((("n3_sim_dreamsim", "DreamSim"), ("n3_sim_dino", "DINO, whole image"))):
        picture(s, name, 0.75 + i * 2.55, 1.25, h=3.4)
    text(s, 5.95, 1.45, 3.7, 3.3,
         [("150 random patches per α, all pairs\n\n", {"size": 13, "color": GREY}),
          ("distance rises with α, peaks at 2.25–2.5\n\n", {"size": 15}),
          ("even the most distinct 5% (0.215) are closer than fractal pairs (0.256)\n\n", {"size": 15}),
          ("dashed red: the same patch shifted 4 px", {"size": 12, "color": GREY})])


def s_sim_other(prs):
    s = new_slide(prs, "Two measures that mislead")
    for i, name in enumerate(("n3_sim_local", "n3_sim_pixel")):
        picture(s, name, 0.75 + i * 2.55, 1.25, h=3.4)
    text(s, 5.95, 1.45, 3.7, 3.3,
         [("DINO, location by location\n", {"size": 15}),
          ("noise looks as different as fractals (0.31–0.34), but a 4 px shift already scores 0.16\n\n",
           {"size": 12, "color": GREY}),
          ("pixels\n", {"size": 15}),
          ("independent patches are uncorrelated at every α; a 4 px shift decorrelates fine noise too",
           {"size": 12, "color": GREY})])


def s_sim_between(prs):
    s = new_slide(prs, "Across α: a different texture")
    picture(s, "n3_between", 0.8, 1.3, h=3.4)
    text(s, 5.2, 1.6, 4.4, 2.8,
         [("between two α values, distances are large\n", {"size": 15}),
          ("α 1 vs 2 ≈ 0.4; within one α 0.02–0.13\n\n", {"size": 12, "color": GREY}),
          ("the networks read noise as “which texture”, not “which patch”\n\n", {"size": 15}),
          ("so α sets difficulty; the pair within α is what is learned", {"size": 12, "color": GREY})])


# ---------------------------------------------------------------- slides: predictions

def prediction_tag(slide):
    pill(slide, 8.05, 0.55, 1.5, 0.36, "prediction", "FFA630", size=12)


def s_pred_baseline(prs):
    s = new_slide(prs, "d′ before learning rises with α")
    prediction_tag(s)
    picture(s, "n3_pred_baseline", 0.5, 1.3, h=3.2)
    text(s, 6.35, 1.55, 3.3, 3.0,
         [("assumed: d′ ∝ DreamSim distance\n\n", {"size": 15}),
          ("study α: 22–49% of a fractal pair's d′\n\n", {"size": 15}),
          ("at 6° from fixation patches likely look even more alike, most at low α",
           {"size": 12, "color": GREY})])


def s_pred_learning(prs):
    s = new_slide(prs, "Learning the rule depends on seeing the patches")
    prediction_tag(s)
    picture(s, "n3_pred_learning", 0.4, 1.2, h=3.95)
    text(s, 6.0, 1.45, 3.7, 3.4,
         [("kernel model + misread patches: each one is taken for its pair partner with p = 1 − Φ(d′/2)\n\n",
           {"size": 12, "color": GREY}),
          ("criterion reached only above d′ ≈ 2.5\n\n", {"size": 15}),
          ("where the study's α fall depends on the fractal d′, which the pilot measures",
           {"size": 13, "color": GREY})])


def s_pred_gain(prs):
    s = new_slide(prs, "Learning should sharpen d′ most in the middle")
    prediction_tag(s)
    picture(s, "n3_pred_gain", 0.4, 1.2, h=3.95)
    text(s, 6.0, 1.45, 3.7, 3.4,
         [("gain = how much was learned × room to improve\n\n", {"size": 12, "color": GREY}),
          ("too hard: nothing learned · too easy: nothing to gain\n\n", {"size": 15}),
          ("C (top) gains less and needs larger d′: no linear mode in this rule, and not trained in practice",
           {"size": 13, "color": GREY})])


def s_posttest(prs):
    s = new_slide(prs, "Measuring d′: a post-test only")
    pill(s, 7.75, 0.55, 1.8, 0.36, "design idea", "4CC9F0", size=12)
    u, cx, cy = 0.17, 2.3, 3.2
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, cx - 9.5 * u, cy - 9 * u, 19 * u, 15 * u, fill=SCREEN, adj=0.05)
    image(s, "n3_a2.00_1", cx - 5.196 * u, cy + 3 * u, 4 * u)
    image(s, "n3_a2.00_2", cx + 5.196 * u, cy + 3 * u, 4 * u)
    bullseye(s, cx, cy, u)
    text(s, cx - 1.6, cy - 8.5 * u, 3.2, 0.35, "same or different?", size=14, color=WHITE, align="ctr")
    text(s, 4.5, 1.4, 5.1, 3.6,
         [("same / different at the 6° positions\n", {"size": 15}),
          ("~100 trials, ~5 min, after the test\n\n", {"size": 12, "color": GREY}),
          ("trained pairs vs untrained pairs matched in DreamSim distance\n", {"size": 15}),
          ("gives d′(α) and the learning gain\n\n", {"size": 12, "color": GREY}),
          ("no pre-test: it could itself change learning\n", {"size": 15}),
          ("add a pre-test arm only if that effect matters", {"size": 12, "color": GREY})])


# ---------------------------------------------------------------- slides: participants and cost
# Same cost model as the 4x4 deck: Prolific (academic) $12/h, 33.3% fee on all payments,
# bonuses included. Typical accuracy 72% is an assumption (K06 predicts less at small alpha).
PILOT_PER_ALPHA = 10
PAY_PER_HOUR = 12.0
PROLIFIC_FEE = 1 / 3
TYPICAL_ACCURACY = 0.72


def s_participants(prs):
    s = new_slide(prs, "Pilot: how many participants?")
    n_alpha = len(CFG["alphas"])
    people = n_alpha * PILOT_PER_ALPHA
    text(s, 0.6, 1.35, 8.8, 0.9, f"{people}", size=54, color=A_COL, bold=True, align="ctr", anchor="ctr")
    text(s, 0.6, 2.3, 8.8, 0.4, f"{n_alpha} α conditions × {PILOT_PER_ALPHA}", size=18, align="ctr")
    for i, (a, col) in enumerate(zip(CFG["alphas"], STUDY_COLORS)):
        pill(s, 2.35 + i * 1.4, 2.95, 1.2, 0.45, f"α {a}", col, size=14)
    text(s, 0.6, 3.85, 8.8, 0.9,
         [("one Prolific study per α (&alpha=… in its link) → exactly 10 each\n", {"size": 14}),
          ("everyone: the same practice, the same test rule; F / J mapping random", {"size": 12, "color": GREY})],
         align="ctr")


def s_cost(prs):
    s = new_slide(prs, "What the pilot costs on Prolific")
    n = sum(CFG["test_trials"])
    bonus = lambda p: n * (p * CFG["correct"] + (1 - p) * CFG["wrong"]) / CFG["coins_per_dollar"]
    base = round(PAY_PER_HOUR * session_minutes() / 60, 2)
    people = len(CFG["alphas"]) * PILOT_PER_ALPHA
    text(s, 0.6, 1.3, 8.8, 0.7,
         [(f"per person: ${base:.2f} base ({session_minutes()} min at ${PAY_PER_HOUR:.0f}/h) + bonus\n",
           {"size": 15}),
          (f"Prolific adds {PROLIFIC_FEE:.1%} to everything paid, bonuses included", {"size": 13, "color": GREY})],
         align="ctr")
    cases = [("everyone guessing", 0.5), (f"typical ({TYPICAL_ACCURACY:.0%})", TYPICAL_ACCURACY),
             ("everyone at criterion", 1.0)]
    for i, (name, p) in enumerate(cases):
        y = 2.3 + i * 0.6
        total = people * (base + bonus(p)) * (1 + PROLIFIC_FEE)
        text(s, 1.6, y, 3.2, 0.5, name, size=17, color=WHITE if i == 1 else GREY, anchor="ctr")
        text(s, 4.8, y, 1.6, 0.5, f"+${bonus(p):.2f}", size=15, color=COIN_COL, anchor="ctr", align="r")
        text(s, 6.6, y, 1.8, 0.5, f"${total:,.0f}", size=20 if i == 1 else 17, color=WHITE if i == 1 else GREY,
             bold=i == 1, anchor="ctr", align="r")
    text(s, 0.6, 4.35, 8.8, 0.7, [(f"{people} people · add ~10% for replacing excluded participants\n", {"size": 12}),
                                  ("the model predicts below 72% at small α, so the typical bonus may be lower",
                                   {"size": 12})], color=GREY, align="ctr")


SLIDES = [s_title, s_machine, s_rule, s_alphas, s_alpha_range, s_noise_how, s_noise_pairs, s_sim_within,
          s_sim_other, s_sim_between, s_round, s_keys, s_feedback, s_bonus, s_practice, s_test, s_criterion,
          s_dva, s_online, s_session, s_pred_baseline, s_pred_learning, s_pred_gain, s_posttest,
          s_participants, s_cost]


def main():
    make_assets()
    prs = Presentation(str(TEMPLATE))
    ids = prs.slides._sldIdLst
    for sld in list(ids):
        prs.part.drop_rel(sld.get(qn("r:id")))
        ids.remove(sld)
    for i, make in enumerate(SLIDES, 1):
        if ONLY is None or str(i) in ONLY.split(","):
            make(prs)
    prs.save(str(OUT))
    print(f"saved {OUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
