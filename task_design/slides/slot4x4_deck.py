"""4x4 slot-machine study deck, in the style of ~/Downloads/2x2 Design Exp0 + next studies (1).pptx.

Describes the online task in catlearn_4x4_prolific: one idea per slide, little text, no
predictions. Timing, coins, practice criterion, test types and rule tables are read from
the task's own src/config.js, so the slides match what participants actually see.

Keeps the template's master, layouts and embedded Nunito font, drops its slides, and draws
everything (machine, rounds, feedback, rule grids, calibration) as editable shapes with the
real fractal images.

  python slot4x4_deck.py [out.pptx] [slide numbers to preview, e.g. 2,3]
"""

import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

import design_deck as dd
from design_deck import A_COL, B_COL, GREY, SCREEN, WHITE, image, label_in, line, shape, text

REPO = Path(__file__).resolve().parents[2]
CONFIG = REPO / "catlearn_4x4_prolific" / "src" / "config.js"
TEMPLATE = Path.home() / "Downloads" / "2x2 Design Exp0 + next studies (1).pptx"
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path.home() / "Downloads" / "4x4 Slot Machine Study.pptx"
ONLY = sys.argv[2] if len(sys.argv) > 2 else None

GOLD, GOLD_DARK, KNOB = "C9A227", "6B4E00", "D62828"
WIN_COL, LOSE_COL = "FFD166", "4A4A4A"          # rule grids: winning / losing pairs
COIN_COL = "FFD166"
# Example fractals: two of the task's own matched, equally distinct test groups
# (stimuli/fractal_groups/groups.js), copied into the slide assets as fg<k>.png
GROUP_DIR = REPO / "catlearn_4x4_prolific" / "stimuli" / "fractal_groups"
FRACT_A = [f"fg{k}" for k in (1, 2, 3, 4)]
FRACT_B = [f"fg{k}" for k in (5, 6, 7, 8)]


def make_group_assets():
    """Cropped, downscaled copies of the example fractals for the slides."""
    import numpy as np
    from PIL import Image
    dd.ASSETS.mkdir(parents=True, exist_ok=True)
    for name in FRACT_A + FRACT_B:
        dst = dd.ASSETS / f"{name}.png"
        if dst.exists():
            continue
        im = Image.open(GROUP_DIR / f"{name[2:]}.png").convert("RGBA")
        a = np.array(im)[..., 3]
        r, c = np.where(a.any(1))[0], np.where(a.any(0))[0]
        im = im.crop((c[0], r[0], c[-1] + 1, r[-1] + 1))
        side = max(im.size)
        canvas = Image.new("RGBA", (side, side), (0, 0, 0, 0))
        canvas.paste(im, ((side - im.width) // 2, (side - im.height) // 2))
        canvas.resize((360, 360), Image.LANCZOS).save(dst)


# ---------------------------------------------------------------- task settings from config.js

def read_config():
    js = CONFIG.read_text()

    def block(name):
        m = re.search(rf"export const {name} = (\{{.*?\n\}});", js, re.S)
        if not m:
            raise ValueError(f"{name} not found in {CONFIG}")
        return m.group(1)

    def num(src, key):
        m = re.search(rf"\b{key}:\s*(-?[\d.]+)", src)
        if not m:
            raise ValueError(f"{key} not found")
        return float(m.group(1))

    timing, reward, crit = block("TIMING"), block("REWARD"), block("PRACTICE_CRITERION")
    rules = {}
    for name, typ, size, labels in re.findall(
            r"(\w+):\s*\{\s*type:\s*'([^']+)',\s*size:\s*(\d+),\s*labels:\s*\[([^\]]+)\]", block("RULES")):
        rules[name] = {"type": typ, "size": int(size), "labels": [int(v) for v in labels.split(",")]}
    test_types = re.findall(r"(\w+):\s*\['(\w+)',\s*'(\w+)'\]", block("TEST_TYPES"))
    blocks = re.search(r"export const BLOCKS = \[(.*?)\n\];", js, re.S).group(1)
    reps = {ph: [int(r) for r in re.findall(rf"phase:\s*'{ph}'[^}}]*?reps:\s*(\d+)", blocks)]
            for ph in ("practice", "test")}
    cfg = {
        "wait": num(timing, "wait") / 1000, "pull": num(timing, "pull") / 1000,
        "response": num(timing, "response") / 1000, "feedback": num(timing, "feedback") / 1000,
        "correct": int(num(reward, "correct")), "wrong": int(num(reward, "wrong")), "late": int(num(reward, "late")),
        "coins_per_dollar": int(num(reward, "coinsPerDollar")),
        "window": int(num(crit, "window")), "min_correct": int(num(crit, "minCorrect")),
        "rules": rules, "test_types": test_types,
        "practice_max": [r * 4 for r in reps["practice"]], "test_trials": [r * 16 for r in reps["test"]],
        "assumed_cm": num(block("CALIBRATION"), "assumedDistanceMm") / 10,
    }
    cfg["dist_range_cm"] = [int(v) / 10 for v in re.search(
        r"plausibleDistanceMm:\s*\[(\d+),\s*(\d+)\]", js).groups()]
    study = block("STUDY")
    cfg["study_types"] = {ph: re.findall(r"'(\w+)'", re.search(rf"{ph}:\s*\[([^\]]*)\]", study).group(1))
                          for ph in ("pilot", "main")}
    return cfg


CFG = read_config()


# ---------------------------------------------------------------- drawing helpers

def new_slide(prs, ttl=None):
    """Slide on the template's TITLE_ONLY (or BLANK) layout, looked up by name: layout order
    differs between template decks."""
    layouts = {l.name: l for l in prs.slide_layouts}
    s = prs.slides.add_slide(layouts["TITLE_ONLY"] if ttl else layouts["BLANK"])
    for ph in list(s.placeholders):
        if ph.placeholder_format.type != 1:          # keep only the title
            ph._element.getparent().remove(ph._element)
    if ttl:
        dd.title(s, ttl)
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = dd.rgb("000000")
    return s


def oval(slide, cx, cy, d, fill=None, line_col=None, line_w=1.0):
    return shape(slide, MSO_SHAPE.OVAL, cx - d / 2, cy - d / 2, d, d, fill=fill, line=line_col, line_w=line_w)


def bullseye(slide, cx, cy, u):
    """MATLAB-task fixation: 0.6 deg black disc, white cross 0.15 deg thick, 0.2 deg black centre."""
    d, bar = 0.6 * u, 0.15 * u
    oval(slide, cx, cy, d, fill="000000")
    shape(slide, MSO_SHAPE.RECTANGLE, cx - d / 2, cy - bar / 2, d, bar, fill=WHITE)
    shape(slide, MSO_SHAPE.RECTANGLE, cx - bar / 2, cy - d / 2, bar, d, fill=WHITE)
    oval(slide, cx, cy, 0.2 * u, fill="000000")


def coin_bar(slide, cx, cy, width, frac, label, size=10):
    h = max(0.05, width * 0.04)
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, cx - width / 2, cy - h / 2, width, h, fill="333333", adj=0.5)
    if frac > 0:
        shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, cx - width / 2, cy - h / 2, width * frac, h, fill=COIN_COL, adj=0.5)
    text(slide, cx + width / 2 + 0.06, cy - 0.15, 1.4, 0.3, label, size=size, color=COIN_COL, bold=True, anchor="ctr")


def slot(slide, cx, cy, u, files=None, outcome=None, pulled=False, bar=None, hint=None, screen=True, labels=False):
    """The task screen, u inches per degree: grey display, gold machine (20 x 8.5 deg) with a
    plate and lever, symbols 4 deg at +-6 deg, bull's eye, optional feedback circles, coin
    bar (frac, label) and hint text. Returns the display's bounding box."""
    sw, sh = 26 * u, 15.5 * u
    box = (cx - sw / 2, cy - sh / 2 - 1.2 * u, sw, sh)
    if screen:
        shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, *box, fill=SCREEN, adj=0.04)
    w, h = 20 * u, 8.5 * u
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, cx - w / 2, cy - h / 2, w, h, line=GOLD, line_w=max(0.75, 0.5 * u * 72), adj=0.1)
    plate = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, cx - 1.6 * u, cy - h / 2 - 0.45 * u, 3.2 * u, 0.9 * u, fill=GOLD, adj=0.3)
    if u > 0.12:
        label_in(plate, "SLOTS", size=max(6, int(u * 18)), color="3A2A00", bold=True)
    px = cx + w / 2 + 0.9 * u
    arm_top = cy if pulled else cy - 3.5 * u
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, px - 0.15 * u, arm_top, 0.3 * u, 3.5 * u, fill=GOLD, adj=0.5)
    oval(slide, px, cy + 3.5 * u if pulled else cy - 3.5 * u, 1.0 * u, fill=KNOB)
    if files:
        for f, x in zip(files, (cx - 6 * u, cx + 6 * u)):
            image(slide, f, x, cy, 4 * u)
    if outcome:
        for x in (cx - 6 * u, cx + 6 * u):
            oval(slide, x, cy, 5.4 * u, line_col=outcome, line_w=max(1.0, 0.3 * u * 72))
    bullseye(slide, cx, cy, u)
    if bar is not None:
        frac, lab = bar
        coin_bar(slide, cx - 1.0 * u, cy - h / 2 - 2.6 * u, 9 * u, frac, lab, size=max(7, int(u * 30)))
    if hint:
        text(slide, cx - 8 * u, cy + h / 2 + 0.4 * u, 16 * u, 0.8 * u, hint, size=max(7, int(u * 26)), color="E0E0E0",
             align="ctr", anchor="ctr")
    if labels:
        text(slide, cx - 6 * u - 1, cy - 2.9 * u - 0.35, 2, 0.3, "A", size=16, color=A_COL, bold=True, align="ctr")
        text(slide, cx + 6 * u - 1, cy - 2.9 * u - 0.35, 2, 0.3, "B", size=16, color=B_COL, bold=True, align="ctr")
    return box


def rule_grid(slide, x, y, cell, labels, size, gap=0.02):
    """Win / lose grid: rows = left (A) symbol, columns = right (B) symbol."""
    for a in range(size):
        for b in range(size):
            win = labels[size * a + b] == 1
            shape(slide, MSO_SHAPE.RECTANGLE, x + b * cell, y + a * cell, cell - gap, cell - gap,
                  fill=WIN_COL if win else LOSE_COL)


def key_cap(slide, x, y, letter, size=0.55, label=None, label_col=WHITE):
    k = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, size, size, fill="1E1E1E", line=WHITE, line_w=1.5, adj=0.15)
    label_in(k, letter, size=int(size * 36), color=WHITE, bold=True)
    if label:
        text(slide, x - 0.3, y + size + 0.05, size + 0.6, 0.3, label, size=13, color=label_col, align="ctr", bold=True)


def pill(slide, x, y, w, h, txt, fill, size=14, color="111111"):
    p = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, adj=0.5)
    label_in(p, txt, size=size, color=color, bold=True)
    return p


def arrow(slide, x1, y, x2):
    line(slide, x1, y, x2, y, color=GREY, width=1.5, arrow=True)


# ---------------------------------------------------------------- slides

def s_title(prs):
    s = prs.slides.add_slide(prs.slide_layouts[0])
    for ph in list(s.placeholders):
        if ph.placeholder_format.type not in (3, 4):
            ph._element.getparent().remove(ph._element)
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = dd.rgb("000000")
    for ph in s.placeholders:
        title = ph.placeholder_format.type == 3
        ph.text = "Slot machines: which pairs win?" if title else "A 4×4 online category-learning study"
        for p in ph.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.name = dd.FONT
                r.font.size = Pt(34 if title else 18)
                r.font.color.rgb = dd.rgb(WHITE)
    for i, f in enumerate(FRACT_A + FRACT_B):
        image(s, f, 1.65 + i * 0.95, 4.55, 0.62)


def s_machine(prs):
    s = new_slide(prs, "The slot machine")
    u = 0.24
    cx, cy = 5.0, 3.45
    slot(s, cx, cy, u, files=[FRACT_A[0], FRACT_B[0]], bar=(0.4, "1,240 coins"), labels=True)
    y = cy + 2.55 * u + 0.18
    line(s, cx, y, cx - 6 * u, y, color=WHITE, width=1.25, arrow=True)
    text(s, cx - 6 * u, y + 0.02, 6 * u, 0.28, "6°", size=13, align="ctr")
    text(s, cx + 6 * u - 1, cy + 2 * u + 0.02, 2, 0.28, "4°", size=13, align="ctr")
    text(s, 0.6, 5.08, 8.8, 0.4, "Two symbols, left (A) and right (B) of a bull's-eye fixation · sizes in degrees",
         size=14, color=GREY, align="ctr")


def s_pairs(prs):
    s = new_slide(prs, "4 symbols per side → 16 pairs")
    cell, x0, y0 = 0.6, 4.0, 1.95
    for i, f in enumerate(FRACT_A):
        image(s, f, x0 - 0.45, y0 + (i + 0.5) * cell, 0.5)
    for j, f in enumerate(FRACT_B):
        image(s, f, x0 + (j + 0.5) * cell, y0 - 0.38, 0.5)
    rule_grid(s, x0, y0, cell, CFG["rules"]["x4_X_A"]["labels"], 4, gap=0.04)
    text(s, x0 - 1.25, y0 + 2 * cell - 0.2, 0.5, 0.4, "A", size=20, color=A_COL, bold=True, align="r")
    text(s, x0 - 0.75, y0 - 0.6, 0.6, 0.4, "B", size=20, color=B_COL, bold=True, align="r")
    lx = 7.2
    shape(s, MSO_SHAPE.RECTANGLE, lx, 2.3, 0.35, 0.35, fill=WIN_COL)
    text(s, lx + 0.45, 2.28, 1.8, 0.4, "wins", size=16, anchor="ctr")
    shape(s, MSO_SHAPE.RECTANGLE, lx, 2.85, 0.35, 0.35, fill=LOSE_COL)
    text(s, lx + 0.45, 2.83, 1.8, 0.4, "loses", size=16, anchor="ctr")
    text(s, 0.6, 4.75, 8.8, 0.4, "A rule says which pairs win · fractals are assigned to rows and columns at random",
         size=14, color=GREY, align="ctr")


def s_round(prs):
    s = new_slide(prs, "One round")
    u = 0.1
    steps = [
        ("Wait", f"SPACE, or {CFG['wait']:g} s", dict(hint="press SPACE to play")),
        ("Lever pull", f"{CFG['pull']:g} s", dict(pulled=True)),
        ("Will this pair win?", f"F / J, ≤ {CFG['response']:g} s", dict(files=[FRACT_A[0], FRACT_B[0]])),
        ("Feedback", f"{CFG['feedback']:g} s", dict(files=[FRACT_A[0], FRACT_B[0]], outcome="2ECC40")),
    ]
    for i, (name, dur, kw) in enumerate(steps):
        cx, cy = 1.9 + i * 1.55, 1.95 + i * 0.85
        slot(s, cx, cy, u, bar=(0.4, ""), **kw)
        text(s, cx + 1.4, cy - 0.75, 2.6, 0.7, [(name + "\n", {"size": 15}), (dur, {"size": 13, "color": GREY})])
    line(s, 0.55, 2.45, 3.6, 5.3, color=WHITE, width=1.5, arrow=True)
    text(s, 6.4, 4.85, 3.4, 0.35, "then the next round", size=13, color=GREY, align="r")


def s_keys(prs):
    s = new_slide(prs, "Will this pair win? YES or NO")
    for i, (yes, no) in enumerate((("F", "J"), ("J", "F"))):
        x = 1.4 + i * 4.2
        text(s, x - 0.2, 1.45, 3.2, 0.35, f"Half the participants", size=14, color=GREY, align="ctr")
        k1, k2 = ("F", "J")
        key_cap(s, x + 0.55, 2.0, k1, size=0.8, label="YES" if yes == "F" else "NO",
                label_col="2ECC40" if yes == "F" else "FF6B6B")
        key_cap(s, x + 1.65, 2.0, k2, size=0.8, label="YES" if yes == "J" else "NO",
                label_col="2ECC40" if yes == "J" else "FF6B6B")
    text(s, 0.6, 3.9, 8.8, 0.9, [("Left vs right hand is not tied to YES / NO\n", {"size": 16}),
                                 ("assigned at random per participant · upper or lower case both count",
                                  {"size": 13, "color": GREY})], align="ctr")


def s_feedback(prs):
    s = new_slide(prs, "Feedback: green or red circles")
    u = 0.15
    for i, (col, lab, sub, frac, bar_lab) in enumerate((
            ("2ECC40", f"+{CFG['correct']} coins", "correct", 0.45, "1,250 coins"),
            ("FF4136", f"−{abs(CFG['wrong'])} coins", "wrong or too slow", 0.42, "1,235 coins"))):
        cx = 2.55 + i * 4.9
        slot(s, cx, 2.95, u, files=[FRACT_A[1], FRACT_B[2]], outcome=col, bar=(frac, bar_lab))
        text(s, cx - 1.8, 4.45, 3.6, 0.6, [(lab + "\n", {"size": 18, "color": col}), (sub, {"size": 13, "color": GREY})],
             align="ctr")
    text(s, 0.6, 5.0, 8.8, 0.35, "no text on the screen · the fixation stays on · the coin bar moves", size=13,
         color=GREY, align="ctr")


def s_bonus(prs):
    s = new_slide(prs, "Coins become a bonus")
    n = sum(CFG["test_trials"])
    coins = lambda p: round(n * (p * CFG["correct"] + (1 - p) * CFG["wrong"]))
    coin_bar(s, 4.6, 1.75, 5.0, 0.55, f"{coins(0.8) // 2:,} coins", size=14)
    rows = [("guessing", 0.5), ("80% correct", 0.8), ("perfect", 1.0)]
    for i, (name, p) in enumerate(rows):
        y = 2.4 + i * 0.6
        text(s, 2.2, y, 2.6, 0.5, name, size=17, anchor="ctr")
        text(s, 4.8, y, 1.8, 0.5, f"{coins(p):,} coins", size=17, color=COIN_COL, anchor="ctr", align="r")
        text(s, 6.8, y, 1.2, 0.5, f"${coins(p) / CFG['coins_per_dollar']:.2f}", size=17, color=GREY, anchor="ctr", align="r")
    text(s, 0.6, 4.4, 8.8, 0.8,
         [(f"Test blocks only · {CFG['coins_per_dollar']:,} coins = $1 · never below $0\n", {"size": 14}),
          (f"shown to participants in coins; paid as a Prolific bonus on top of the base pay ({n} test rounds)",
           {"size": 12, "color": GREY})], align="ctr")


def s_practice(prs):
    s = new_slide(prs, "Practice: two 2×2 machines")
    r = CFG["rules"]
    items = [("Type I", "A or B, at random", r["x2_I_A"]["labels"]), ("XOR", "both sides matter", r["x2_XOR"]["labels"])]
    for i, (name, sub, labels) in enumerate(items):
        x = 2.0 + i * 3.6
        pill(s, x, 1.45, 2.2, 0.5, name, A_COL if i == 0 else "3CC9A0", size=17)
        rule_grid(s, x + 0.55, 2.2, 0.55, labels, 2, gap=0.04)
        text(s, x - 0.2, 3.4, 2.6, 0.35, sub, size=13, color=GREY, align="ctr")
        if i == 0:
            arrow(s, x + 2.35, 1.7, x + 3.45)
    text(s, 0.6, 4.1, 8.8, 0.9,
         [(f"Each ends once {CFG['min_correct']} of the last {CFG['window']} answers are right\n", {"size": 17}),
          (f"at most {CFG['practice_max'][0]} rounds · practice coins are not paid", {"size": 13, "color": GREY})],
         align="ctr")


def s_types(prs):
    s = new_slide(prs, "Test: one rule type per person")
    types = CFG["test_types"]
    cell = 0.24
    for i, (typ, rule_a, _) in enumerate(types):
        x = 0.55 + i * 1.3
        text(s, x - 0.1, 1.55, 1.2, 0.35, typ, size=16, bold=True, align="ctr")
        rule_grid(s, x + 0.02, 2.0, cell, CFG["rules"][rule_a]["labels"], 4)
    text(s, 0.6, 3.25, 8.8, 0.35, "← more of the left side alone            more of the interaction (XOR) →", size=13,
         color=GREY, align="ctr")
    text(s, 0.6, 4.2, 8.8, 0.8, [("Assigned at random · rows = left symbol, columns = right symbol\n", {"size": 15}),
                                 ("A versions shown; type XVI is pure XOR and has no A / B version",
                                  {"size": 12, "color": GREY})], align="ctr")


def s_blocks(prs):
    s = new_slide(prs, "Three test blocks: A, A, then B")
    r = CFG["rules"]
    for b in range(3):
        text(s, 2.4 + b * 1.9, 1.35, 1.5, 0.3, f"block {b + 1}", size=13, color=GREY, align="ctr")
    for row, order in enumerate((("A", "A", "B"), ("B", "B", "A"))):
        y = 1.8 + row * 0.75
        text(s, 0.4, y, 1.9, 0.5, "half" if row == 0 else "other half", size=13, color=GREY, anchor="ctr", align="r")
        for b, v in enumerate(order):
            x = 2.4 + b * 1.9
            pill(s, x, y, 1.5, 0.5, f"{v} version", A_COL if v == "A" else B_COL, size=14)
            if b < 2:
                arrow(s, x + 1.56, y + 0.25, x + 1.84)
    for i, (rule, lab) in enumerate((("x4_X_A", "A version (type X)"), ("x4_X_B", "B version"))):
        x = 3.3 + i * 2.0
        rule_grid(s, x, 3.45, 0.2, r[rule]["labels"], 4)
        text(s, x - 0.45, 4.3, 1.7, 0.3, lab, size=11, color=GREY, align="ctr")
    text(s, 0.6, 4.75, 8.8, 0.35, f"same type, dominant side switches · new fractals every block · "
         f"{CFG['test_trials'][0]} rounds per block", size=13, color=GREY, align="ctr")


def s_dva(prs):
    s = new_slide(prs, "Sizes in degrees of visual angle")
    # card step
    text(s, 0.6, 1.3, 4.2, 0.35, "1 · match a real card", size=16, align="ctr")
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 1.55, 1.85, 2.2, 1.39, fill="DDDDDD", adj=0.06)
    shape(s, MSO_SHAPE.RECTANGLE, 3.5, 3.0, 0.25, 0.24, line="FF0000", line_w=2.5)
    text(s, 0.6, 3.35, 4.2, 0.35, "→ pixels per mm", size=14, color=GREY, align="ctr")
    # blind-spot step
    text(s, 5.2, 1.3, 4.2, 0.35, "2 · find the blind spot", size=16, align="ctr")
    shape(s, MSO_SHAPE.RECTANGLE, 8.35, 2.35, 0.35, 0.35, fill="000000", line=WHITE, line_w=2)
    oval(s, 6.6, 2.52, 0.33, fill="FF2222", line_col=WHITE, line_w=1.5)
    line(s, 6.35, 2.95, 5.75, 2.95, color=GREY, width=1.5, arrow=True)
    text(s, 5.2, 3.05, 4.2, 0.6, [("cover the right eye, look AT the square,\n", {"size": 12}),
                                  ("SPACE when the ball vanishes → viewing distance", {"size": 12})],
         color=GREY, align="ctr")
    text(s, 0.6, 3.95, 8.8, 0.45, "pixels per degree = 2 · distance · tan(0.5°) · pixels per mm", size=16,
         align="ctr")
    lo, hi = CFG["dist_range_cm"]
    text(s, 0.6, 4.45, 8.8, 0.7,
         [(f"Outside {lo:g}–{hi:g} cm: measure again once, then assume {CFG['assumed_cm']:g} cm (flagged)\n", {}),
          ("too big for the window: everything shrinks by one factor (layout_scale, saved)", {})],
         size=12, color=GREY, align="ctr")


def s_online(prs):
    s = new_slide(prs, "Built for online data")
    items = [
        ("Desktop only", "browser check, window ≥ 1000 × 650 px"),
        ("Full screen", "exits are counted on every trial"),
        ("Tab switches", "counted on every trial"),
        ("Quiz", "instructions repeat until right (3 tries)"),
        ("Every trial saved", "pair, key, RT, coins, px / degree, how the round started"),
        ("Reproducible", "seed rebuilds the design · Prolific IDs on every row"),
    ]
    for i, (head, sub) in enumerate(items):
        col, row = i % 2, i // 2
        x, y = 0.9 + col * 4.4, 1.5 + row * 1.1
        shape(s, MSO_SHAPE.OVAL, x, y + 0.08, 0.22, 0.22, fill=A_COL)
        text(s, x + 0.4, y - 0.05, 3.9, 0.9, [(head + "\n", {"size": 16}), (sub, {"size": 12, "color": GREY})])


def session_minutes():
    """Rough duration: ~0.8 s wait + pull + ~1.2 s answer + feedback per round, ~40 practice
    rounds per practice block, ~8 min for consent, screen setup and instructions."""
    per_round = 0.8 + CFG["pull"] + 1.2 + CFG["feedback"]
    return round((sum(CFG["test_trials"]) * per_round + 2 * 40 * per_round) / 60 + 8)


def s_session(prs):
    s = new_slide(prs, "The session")
    steps = [("Consent", None), ("Screen setup", "~2 min"), ("Instructions + quiz", None),
             ("Practice", "two 2×2 machines"), ("Test", f"3 × {CFG['test_trials'][0]}"), ("Bonus → Prolific", None)]
    x = 0.45
    widths = [1.15, 1.45, 1.75, 1.15, 1.15, 1.6]
    for (name, sub), w in zip(steps, widths):
        pill(s, x, 2.15, w, 0.6, name, "3A3A3A", size=12, color=WHITE)
        if sub:
            text(s, x, 2.85, w, 0.3, sub, size=11, color=GREY, align="ctr")
        x += w + 0.13
    text(s, 0.6, 3.9, 8.8, 0.5, f"about {session_minutes()} minutes", size=20, align="ctr")
    text(s, 0.6, 4.45, 8.8, 0.35, "self-paced rounds · rest between blocks", size=13, color=GREY, align="ctr")


# ---------------------------------------------------------------- the fractals
# How the stimulus set was made (fractal_stimuli/, see its README; numbers and the distance
# histogram from catlearn_4x4_prolific/stimuli/fractal_groups/report.png). OLD_FRACT: the
# earlier catlearn_task set, for comparison.
OLD_FRACT = ["fractal34", "fractal58", "fractal17", "fractal66"]


def lstar_grey(L):
    """sRGB hex of the neutral grey with CIE lightness L*."""
    y = ((L + 16) / 116) ** 3 if L > 8 else L / 903.3
    v = 12.92 * y if y <= 0.0031308 else 1.055 * y ** (1 / 2.4) - 0.055
    g = round(255 * min(1, max(0, v)))
    return f"{g:02X}" * 3


def fractal_row(slide, names, x, y, size, gap):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x - 0.15, y - 0.15, len(names) * (size + gap) - gap + 0.3,
          size + 0.3, fill=SCREEN, adj=0.08)
    for i, f in enumerate(names):
        image(slide, f, x + i * (size + gap) + size / 2, y + size / 2, size)


def s_fractals_why(prs):
    s = new_slide(prs, "New fractals: matched in every way but shape")
    rows = [("before", OLD_FRACT, "brightness L* 35–93 · colourfulness 60–129 · size 15–33%"),
            ("now", FRACT_A, "brightness L* 55–64 · colourfulness 36–40 · size 22% for all")]
    for i, (lab, names, stats) in enumerate(rows):
        y = 1.55 + i * 1.5
        text(s, 0.4, y + 0.3, 1.3, 0.4, lab, size=16, color=GREY if i == 0 else WHITE, align="r")
        fractal_row(s, names, 2.0, y, 0.85, 0.25)
        text(s, 6.6, y + 0.1, 3.2, 0.7, stats.replace(" · ", "\n"), size=12, color=GREY if i == 0 else WHITE)
    text(s, 0.6, 4.75, 8.8, 0.4, "the old set used fully saturated colours: some fractals were much brighter, "
         "more colourful or bigger than others", size=12, color=GREY, align="ctr")


def s_fractals_how(prs):
    s = new_slide(prs, "How each fractal is made")
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.7, 1.45, 3.0, 3.0, fill=SCREEN, adj=0.06)
    image(s, FRACT_A[1], 2.2, 2.95, 2.4)
    text(s, 0.7, 4.55, 3.0, 0.3, "1,500 candidates", size=12, color=GREY, align="ctr")
    x = 4.3
    text(s, x, 1.4, 5.3, 0.6, [("3 nested polygons with bent edges\n", {"size": 16}),
                               ("after Miyashita et al. (1997)", {"size": 12, "color": GREY})])
    text(s, x, 2.15, 5.3, 0.35, "same lightness in every fractal (CIE L*)", size=15)
    for i, (lab, L) in enumerate((("outer", 62), ("middle", 48), ("inner", 68), ("background", 53.6))):
        bx = x + i * 1.3
        sw = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, bx, 2.6, 1.1, 0.5, fill=lstar_grey(L), adj=0.15,
                   line=WHITE if lab == "background" else None, line_w=1)
        label_in(sw, f"{L:g}", size=13, color="111111", bold=True)
        text(s, bx - 0.1, 3.12, 1.3, 0.3, lab, size=11, color=GREY, align="ctr")
    text(s, x, 3.6, 5.3, 0.9, [("3 of 12 evenly spaced hues · chroma 40\n", {"size": 15}),
                               ("lowered to the screen gamut where needed (min ~29, cyan–blue)\n", {"size": 11, "color": GREY}),
                               ("same filled area: 22% of the image", {"size": 15})])


def s_fractals_distinct(prs):
    s = new_slide(prs, "Equally distinct groups")
    fractal_row(s, FRACT_A, 0.85, 1.55, 0.72, 0.2)
    text(s, 0.6, 2.5, 3.9, 0.7, [("every pair in a group: 0.256 ± 0.01\n", {"size": 15}),
                                 ("so every block is equally hard", {"size": 12, "color": GREY})], align="ctr")
    text(s, 0.6, 3.35, 3.9, 1.2, [("distance = DreamSim\n", {"size": 14}),
                                  ("DINO + CLIP + OpenCLIP, tuned on human similarity judgements; "
                                   "1 − cosine, on the task's grey background", {"size": 11, "color": GREY})],
         align="ctr")
    pic = s.shapes.add_picture(str(dd.ASSETS / "fractal_distance_hist.png"), 0, dd.Inches(1.55),
                               height=dd.Inches(1.85))
    pic.left = dd.Inches(4.75 + 4.9 / 2) - pic.width // 2
    text(s, 4.75, 3.5, 4.9, 0.8, [("no two of the 96 chosen closer than 0.162: no near-duplicates\n", {"size": 12}),
                                  ("all candidate pairs: median 0.256, 5–95% range 0.162–0.357", {"size": 11,
                                                                                               "color": GREY})],
         align="ctr")


def s_fractals_set(prs):
    s = new_slide(prs, "The final set")
    s.shapes.add_picture(str(dd.ASSETS / "fractal_contact_sheet.png"), dd.Inches(0.9), dd.Inches(1.2),
                         height=dd.Inches(3.55))
    text(s, 0.6, 4.85, 8.8, 0.4, "18 groups of 4 (test) · 12 pairs (practice) · each side of a block shows one "
         "whole group, never reused for that participant", size=12, color=GREY, align="ctr")


# ---------------------------------------------------------------- participants and cost
# Participants per rule type and which types each phase uses (STUDY in config.js: the pilot
# uses 4 types, the main study all 7), and Prolific pricing for academic accounts
# (researcher-help.prolific.com, Oct 2026): 33.3% fee on all participant payments, bonuses
# included; recommended pay $12/h. Typical accuracy: the adaptive kernel model's mean over
# the test blocks (kernel_model/K04_slot4x4.py), an assumption until the pilot data are in.
SAMPLES = {"Pilot": ("pilot", 10), "Main study": ("main", 30)}
PAY_PER_HOUR = 12.0
PROLIFIC_FEE = 1 / 3
TYPICAL_ACCURACY = 0.72


def per_person():
    n_paid = sum(CFG["test_trials"])
    coins = lambda p: n_paid * (p * CFG["correct"] + (1 - p) * CFG["wrong"])
    base = round(PAY_PER_HOUR * session_minutes() / 60, 2)
    return base, coins(TYPICAL_ACCURACY) / CFG["coins_per_dollar"], coins(1.0) / CFG["coins_per_dollar"]


def sample(name):
    phase, n = SAMPLES[name]
    types = CFG["study_types"][phase]
    return types, n, len(types) * n


def s_participants(prs):
    s = new_slide(prs, "How many participants?")
    for i, name in enumerate(SAMPLES):
        types, n, people = sample(name)
        x = 1.0 + i * 4.4
        text(s, x, 1.3, 3.6, 0.4, name, size=18, color=GREY, align="ctr")
        text(s, x, 1.7, 3.6, 0.85, f"{people}", size=46, color=A_COL if i == 0 else WHITE, bold=True,
             align="ctr", anchor="ctr")
        text(s, x, 2.6, 3.6, 0.4, f"{len(types)} rule types × {n}", size=16, align="ctr")
        text(s, x, 3.0, 3.6, 0.35, ", ".join(types), size=13, color=A_COL if i == 0 else WHITE, align="ctr")
        text(s, x, 3.4, 3.6, 0.6, [(f"within each type: {n // 2} F = YES · {n // 2} J = YES\n", {}),
                                   (f"{n // 2} A first · {n // 2} B first", {})], size=12, color=GREY, align="ctr")
    text(s, 0.6, 4.45, 8.8, 0.7,
         [("pilot: types from opposite ends of the spectrum (VI, X: more one side · XV, XVI: more XOR)\n",
           {"size": 13}),
          ("one Prolific study per type (?type=… in its link), so each gets exactly its places",
           {"size": 12, "color": GREY})], align="ctr")


def s_cost(prs):
    s = new_slide(prs, "What it costs on Prolific")
    base, bonus_typ, bonus_max = per_person()
    fee = PROLIFIC_FEE
    text(s, 0.6, 1.3, 8.8, 0.7,
         [(f"per person: ${base:.2f} base ({session_minutes()} min at ${PAY_PER_HOUR:.0f}/h)"
           f" + bonus ~${bonus_typ:.2f} (up to ${bonus_max:.2f})\n", {"size": 15}),
          (f"Prolific adds {fee:.1%} to everything paid, bonuses included", {"size": 13, "color": GREY})], align="ctr")
    cols = [("", 1.2), ("people", 1.4), ("typical", 1.9), ("maximum", 1.9)]
    x0, y0 = 1.3, 2.35
    x = x0
    for head, w in cols:
        text(s, x, y0, w, 0.35, head, size=13, color=GREY, align="ctr")
        x += w
    for r, name in enumerate(SAMPLES):
        _, _, people = sample(name)
        y = y0 + 0.45 + r * 0.6
        typ = people * (base + bonus_typ) * (1 + fee)
        mx = people * (base + bonus_max) * (1 + fee)
        vals = [name, f"{people}", f"${typ:,.0f}", f"${mx:,.0f}"]
        x = x0
        for (head, w), v in zip(cols, vals):
            text(s, x, y, w, 0.45, v, size=18, color=WHITE if head else GREY, align="ctr", anchor="ctr",
                 bold=head in ("typical",))
            x += w
    text(s, 0.6, 4.25, 8.8, 0.8,
         [(f"typical = ~{TYPICAL_ACCURACY:.0%} correct (kernel-model average) · maximum = every test round correct\n",
           {"size": 12}),
          ("add ~10% for replacing excluded participants · VAT on the fee where it applies", {"size": 12})],
         color=GREY, align="ctr")


SLIDES = [s_title, s_machine, s_pairs, s_fractals_why, s_fractals_how, s_fractals_distinct, s_fractals_set,
          s_round, s_keys, s_feedback, s_bonus, s_practice, s_types, s_blocks, s_dva, s_online, s_session,
          s_participants, s_cost]


def main():
    make_group_assets()
    prs = Presentation(str(TEMPLATE))
    ids = prs.slides._sldIdLst
    for sld in list(ids):
        prs.part.drop_rel(sld.get(qn("r:id")))
        ids.remove(sld)
    for i, make in enumerate(SLIDES, 1):
        if ONLY is None or str(i) in ONLY.split(","):
            make(prs)
    prs.save(str(OUT))
    print(f"saved {OUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
